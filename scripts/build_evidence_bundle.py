#!/usr/bin/env python3
"""Build a Chinese evidence index and embedded image bundle.

Manifest v2 is deliberately evidence-bounded. Version 1 manifests remain
readable, but are always rendered as an internally marked draft because their
boolean checks cannot identify the material supporting each conclusion.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import math
import os
import re
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from path_safety import first_link_component, is_link_or_junction
from validate_provenance import validate_office_provenance


ALLOWED_IMAGES = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff"}
ALLOWED_DISPLAY_ROTATIONS = {0, 90, 180, 270}
CHECK_KEYS = (
    "identity_checked",
    "dates_checked",
    "amounts_checked",
    "payment_status_checked",
    "proof_scope_checked",
    "original_carrier_checked",
)
CHECK_STATUSES = {"resolved", "unresolved", "not_applicable"}
PURPOSE_BASIS_KINDS = {"pleading", "user_statement", "materials_only"}
PURPOSE_BASIS_STATUSES = {"confirmed", "provisional"}
FACT_LEVELS = {"element_fact", "indirect_fact", "auxiliary_fact", "procedural_fact"}
CLAIM_ROLES = {"direct", "indirect", "corroborative", "rebuttal", "linking"}
PURPOSE_TARGET_KINDS = {"claim", "defense", "rebuttal", "context"}
INTAKE_STATUSES = {"not_required", "pending", "completed", "waived_by_user"}
INTAKE_CHOICES = {"provided", "material_review", "unknown", "declined", "later"}
INTAKE_PROTOCOLS = {"field_binding_v1"}
ANSWER_CONTEXT_KINDS = {"free_text", "open", "binary", "options", "summary"}
INTAKE_LOCATOR_KINDS = {"conversation_turn", "transcript_turn"}
BINDING_RESOLUTION_STATUSES = {
    "resolved",
    "ambiguous",
    "unknown",
    "declined",
    "later",
}
BINDING_BASIS_KINDS = {
    "direct_answer",
    "selected_option",
    "material_lookup",
    "computed",
    "confirmed_summary",
}
FACT_VERIFICATION_NOT_VERIFIED = "not_verified_by_intake"
INTAKE_CATEGORIES = {
    "案件信息",
    "提交方",
    "书状",
    "主体",
    "日期",
    "金额",
    "付款状态",
    "上下文",
    "证明用途",
    "原始载体",
    "可读性",
    "其他",
}
INTAKE_CATEGORY_PATTERNS = {
    "案件信息": r"案件信息|案件名称|案号|案由|法院|仲裁机构|审理阶段",
    "提交方": r"提交方|提交身份|以原告身份提交|以被告身份提交|申请人提交|被申请人提交|代理人提交",
    "书状": r"书状|起诉状|答辩状|仲裁申请书|仲裁答辩书|代理意见",
    "主体": r"主体|身份|姓名|名称|账号|当事人|统一社会信用代码",
    "日期": r"日期|时间|年份|年月日",
    "金额": r"金额|数额|数值|总额|余额|主张额|应付额|尚欠|款项|价款|账款|人民币|\d+元",
    "付款状态": r"付款状态|付款情况|支付情况|返还情况|清偿状态|结算状态|结清|欠付|付款|支付|到账|清偿|退款|未付",
    "上下文": r"上下文|前后文|完整聊天|缺页|连续记录|材料完整性",
    "证明用途": r"证明用途|证明对象|待证事实|证明要件|诉讼请求|仲裁请求|抗辩主张|请求权",
    "原始载体": r"原始载体|原手机|原应用|录屏|原始回单|纸质原件|原文件",
    "可读性": r"可读性|清晰|可读|模糊|损坏|方向",
    "其他": r".",
}
RISK_PATTERNS = (
    ("胜诉", re.compile(r"胜诉")),
    ("当然连带责任", re.compile(r"当然连带责任")),
    ("应承担连带责任", re.compile(r"(?:应当?|应|必须|当然)承担连带责任")),
    ("连带责任成立", re.compile(r"连带责任(?:已经?|已)?成立")),
    ("责任已确定", re.compile(r"责任(?:已经?|已)(?:确定|认定)")),
    ("承担责任", re.compile(r"承担责任")),
    ("已实际支付/已清偿", re.compile(r"(?:已|已经)(?:实际支付|实付|清偿|付款|支付|缴纳|交付)")),
    ("判决已生效", re.compile(r"判决(?:已经?|已)(?:生效|发生法律效力)")),
    ("执行结案", re.compile(r"(?:执行|案件)(?:已经?|已)?结案")),
    ("材料直接证明承认", re.compile(r"(?:材料)?直接证明\s*承认")),
    ("承认/自认越界", re.compile(r"(?:承认|自认)")),
    ("直接证明身份真实性", re.compile(r"直接证明.{0,8}(?:身份|签章|真实性)")),
    ("昵称/头像直接锁定身份", re.compile(r"(?:昵称|头像|账号|用户名).{0,12}(?:即|就是|为|对应|足以证明).{0,8}(?:本人|被告|原告|身份|主体)")),
    ("群名/群聊名称直接等同主体", re.compile(r"(?:群名|群聊名称|群名称).{0,16}(?:即|就是|等同于|为|代表|对应|足以证明).{0,16}(?:主体|公司|被告|原告|本人)")),
    ("聊天承认/自认", re.compile(r"(?:聊天|截图|对话|消息|记录).{0,12}(?:承认|自认)")),
    ("另案/关联判决越界", re.compile(r"(?:另案|关联|相关|其他|生效)?判决.{0,20}(?:证明|确定|当然|本案)")),
    ("申请材料等同付款", re.compile(r"(?:申请书|申请材料|申请).{0,12}(?:即|已|证明|代表|视为).{0,8}(?:付款|支付|缴纳|交付)")),
    ("申请材料推定付款", re.compile(r"(?:申请书|申请材料|申请).{0,30}(?:付款|支付|缴纳|交付)")),
    ("申请额等同实际支付额", re.compile(r"(?:申请额|申请金额|requested_amount).{0,18}(?:就是|等于|即为|即是|代表).{0,18}(?:实际支付|实付|实际付款)")),
    ("会还等同债务承认", re.compile(r"(?:会还|会偿还|会还款|会归还)[^。；\n]{0,20}(?:故|所以|因此|即|就是|等于|意味着|表明)[^。；\n]{0,16}(?:债务|欠款)(?:承认|成立|存在)?")),
    ("会还推定债务或欠款成立", re.compile(r"(?:会还|会偿还|会还款|会归还)[^。；\n]{0,18}(?:债务|欠款)(?:承认|成立|存在)")),
    ("债务或欠款成立/存在越界", re.compile(r"(?:债务|欠款).{0,8}(?:成立|存在)")),
    ("材料直接证明合同/违约/债务", re.compile(r"(?:材料|截图|证据|图片|聊天记录).{0,16}(?:证明|足以证明|直接证明).{0,16}(?:合同关系成立|合同成立|被告已违约|债务成立|债务存在|欠款成立|欠款存在)")),
    ("账号/账户直接等同当事人", re.compile(r"(?:账号|账户|微信号|昵称|头像|用户名).{0,12}(?:是|系|即|就是|属于|为|对应|代表).{0,10}(?:被告|原告|申请人|被申请人|本人|责任主体)")),
    ("责任承担终局结论", re.compile(r"(?:被告|原告|申请人|被申请人|第三人)?(?:应当?|必须|当然)?承担(?:连带|补充|赔偿|违约)?责任")),
    ("关联裁判等同本案事实", re.compile(r"(?:另案|关联|相关|其他)?(?:裁判|裁定|判决|裁判文书).{0,18}(?:已|已经)?(?:确定|认定|查明).{0,12}(?:本案)?(?:事实|责任|法律关系)")),
    ("本案事实由关联裁判确定", re.compile(r"(?:本案)?(?:事实|责任|法律关系).{0,12}(?:已|已经)?(?:由|被).{0,8}(?:另案|关联|相关|其他)(?:裁判|裁定|判决|裁判文书).{0,8}(?:确定|认定|查明)")),
    ("合同/违约/债务成立结论", re.compile(r"(?:拟)?证明.{0,12}(?:合同关系|合同|违约|债务|欠款|责任)(?:已经?|已)?(?:成立|存在)|(?:合同关系|合同|违约|债务|欠款|责任)(?:已经?|已)(?:成立|存在)")),
    ("请求应予支持终局结论", re.compile(r"(?:本案)?(?:诉讼|仲裁)?(?:请求|主张).{0,8}(?:应予|应当|应)?(?:支持|成立)")),
)

# Keep filing-facing proof objects affirmative and concise.  These markers are
# display-only: the original text, proof_claims[].boundary and validation data
# remain untouched in the Manifest/QA sidecar and continue to drive all gates.
PROOF_DISPLAY_BOUNDARY_MARKERS = re.compile(
    r"(?:仍|尚)?(?:待核|待核验|待核对|需核验|需核对|需补强|仍需补强|尚需补强)|"
    r"(?:不能|无法|未能|不足以|不得)(?:单独|直接|据此)?(?:证明|认定|推出)|"
    r"不(?:应)?据此(?:证明|认定|推出)|"
    r"(?:原件|真实性|主体身份|签署|送达|清偿状态|付款状态)[^；;。！？!?\n]{0,24}"
    r"(?:待补|待核|待核验|待核对|需补强)"
)
PROOF_DISPLAY_POSITIVE_HINTS = re.compile(
    r"拟证明|用于支持|用于反驳|证明要件|载明|显示|呈现|形成|签署|履行|支付|付款|"
    r"尚未付款|尚未支付|未清偿|催告|通知|交付|债权|价款|金额|日期|主体"
)
CASE_INFO_PLACEHOLDER = re.compile(
    r"(?:案号|案件编号)\s*[：:]?\s*(?:待补|暂缺|缺失|待核|待核验|待核对|未提供|待提供)|"
    r"案件信息\s*[：:]?\s*(?:待补|暂缺|缺失|待核|待核验|待核对|未提供|待提供)"
)
PROOF_OBJECT_TARGET_HINTS = re.compile(
    r"诉讼请求|仲裁请求|请求权|请求第|"
    r"(?:第\s*[一二三四五六七八九十\d]+\s*项|两项|各项)[^；;。！？!?\n]{0,16}请求|"
    r"(?:起诉状|答辩状|仲裁申请书|仲裁答辩书)[^；;。！？!?\n]{0,30}(?:主张|请求|抗辩)|"
    r"抗辩|反驳|法律要件|证明要件|"
    r"构成要件|争议焦点|背景用途|背景事实|时间线背景|程序事项"
)
TARGET_CLAUSE_BOILERPLATE = re.compile(
    r"对应|用于|支持|反驳|说明|明确|关于|相关|服务于|本案|本测试|测试|合成|"
    r"诉讼|仲裁|请求权|请求|抗辩|主张|法律|证明|事实|构成|要件|争议焦点|"
    r"背景用途|背景事实|时间线背景|程序事项|事项|作用|目的|中的|之中的|的|"
    r"(?:起诉状|答辩状|仲裁申请书|仲裁答辩书)|"
    r"第\s*[一二三四五六七八九十百\d]+\s*项"
)
PLEADING_LOCATOR = re.compile(
    r"第\s*[一二三四五六七八九十百\d]+\s*(?:页|段|项|条|款|节|章)|"
    r"页码\s*[：:]?\s*\d+|"
    r"(?:事实与理由|诉讼请求|仲裁请求|答辩理由|请求事项|答辩意见)(?:部分|章节)?"
)
PROOF_OBJECT_MATERIAL_SUMMARY = re.compile(
    r"^(?:拟)?证明(?:该|本|上述)?(?:合同|协议|聊天记录|微信聊天记录|对话记录|转账记录|"
    r"付款记录|银行回单|截图|图片|材料|文件|证据)(?:的)?(?:内容|记载|存在|真实性)?$"
)


class ManifestError(ValueError):
    """Validation error with separate fatal and formal-blocking messages."""

    def __init__(self, errors: Iterable[str], warnings: Iterable[str] = ()):
        self.errors = list(dict.fromkeys(str(item) for item in errors if str(item)))
        self.warnings = list(dict.fromkeys(str(item) for item in warnings if str(item)))
        message = "\n".join(f"- {item}" for item in self.errors) or "清单校验失败"
        if self.warnings:
            message += "\n警示：\n" + "\n".join(f"- {item}" for item in self.warnings)
        super().__init__(message)


def resolve_material_path(raw_path: Any, base_dir: Path) -> Path:
    """Resolve a manifest path without rejecting Chinese or spaced paths."""

    if not isinstance(raw_path, (str, os.PathLike)) or not str(raw_path).strip():
        raise ValueError("材料路径必须是非空字符串")
    raw = os.path.expandvars(str(raw_path))
    if os.name != "nt" and (re.match(r"^[A-Za-z]:[\\/]", raw) or raw.startswith("\\\\")):
        raise ValueError("不得在当前平台使用 Windows 绝对路径或 UNC 路径")
    path = Path(raw).expanduser()
    if not path.is_absolute():
        path = base_dir / path
    if first_link_component(path) is not None:
        raise ValueError("材料路径不得经过符号链接或目录联接")
    resolved = path.resolve()
    return resolved


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _nonempty(value: Any) -> bool:
    if isinstance(value, str):
        return bool(value.strip())
    return value is not None


def _display_proof_object(value: Any) -> str:
    """Remove verification/boundary clauses from the Word table only.

    A defensive tail after a Chinese/ASCII comma may share one sentence with a
    useful proposition, so retain the positive prefix when it is identifiable.
    Substantive negative facts such as “尚未付款” do not match the boundary
    markers and remain visible.
    """

    raw = _text(value).strip()
    if not raw:
        return ""
    kept: list[str] = []
    clauses = re.split(r"(?<=[；;。！？!?])|\n+", raw)
    for clause in clauses:
        clause = clause.strip()
        if not clause:
            continue
        marker = PROOF_DISPLAY_BOUNDARY_MARKERS.search(clause)
        if marker is None:
            kept.append(clause)
            continue
        separator = max(clause.rfind("，", 0, marker.start()), clause.rfind(",", 0, marker.start()))
        prefix = clause[:separator].strip() if separator >= 0 else ""
        if prefix and PROOF_DISPLAY_POSITIVE_HINTS.search(prefix):
            kept.append(prefix.rstrip("；;。！？!?，,") + "。")
    return "".join(kept).strip()


def _display_case_info(value: Any) -> str:
    """Hide missing-information placeholders while preserving known details."""

    text = _text(value).strip()
    if not text:
        return ""
    text = re.sub(r"[；;，,、]?\s*" + CASE_INFO_PLACEHOLDER.pattern, "", text)
    text = re.sub(r"[（(]\s*[；;，,、]*\s*[）)]", "", text)
    text = re.sub(r"([（(])\s*[；;，,、]+", r"\1", text)
    text = re.sub(r"[；;，,、]+\s*([）)])", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text)
    text = re.sub(r"^案件信息\s*[：:]\s*", "", text)
    text = text.strip(" ；;，,、")
    return "" if CASE_INFO_PLACEHOLDER.fullmatch(text) else text


def _purpose_semantic_core(value: Any) -> str:
    """Reduce a purpose to its factual core for duplicate-content checks."""

    text = _text(value).lower()
    for phrase in (
        "提交该材料是为了",
        "提交材料是为了",
        "提交该材料",
        "提交材料",
        "拟用于",
        "用于",
        "支持",
        "反驳",
        "本案",
        "服务于",
        "拟证明",
        "证明",
        "请求",
        "抗辩",
        "争议焦点",
        "证明要件",
        "事实要件",
        "构成要件",
        "法律要件",
        "事实",
        "要件",
    ):
        text = text.replace(phrase, "")
    text = re.sub(r"\b(?:m|s|r|ctx)[-_]?\d{1,4}\b", "", text, flags=re.IGNORECASE)
    return re.sub(r"[\s。；，：:、,.，;！？!?（）()\[\]{}]+", "", text).replace("的", "").replace("该", "")


def _purpose_has_concrete_anchor(value: Any) -> bool:
    """Identify a non-generic request/defense anchor before duplicate gating."""

    text = _text(value)
    token_anchor = any(
        token in text
        for token in (
            "返还",
            "赔偿",
            "清偿",
            "履行",
            "解除",
            "违约",
            "损失",
            "价款",
            "借款",
            "劳务",
            "工资",
            "律师费",
            "代理费",
            "转让",
            "股权",
            "支付",
            "交付",
            "驳回",
            "撤销",
            "无效",
            "起诉状第",
            "答辩状第",
            "仲裁请求第",
            "仲裁答辩",
        )
    )
    confirmation_anchor = bool(
        re.search(r"确认(?:合同|关系|权利|义务|事实|效力|无效|所有权|债权|债务)", text)
    )
    return token_anchor or confirmation_anchor


def _target_clause_has_specific_content(value: Any, *, display: bool = False) -> bool:
    """Reject a target clause made only from generic litigation labels."""

    text = _display_proof_object(value) if display else _text(value).strip()
    for clause in re.split(r"[；;。！？!?\n]+", text):
        if not (
            PROOF_OBJECT_TARGET_HINTS.search(clause)
            or re.search(r"请求|抗辩|要件|争议|主张|反驳|诉讼|仲裁|背景|程序", clause)
        ):
            continue
        remainder = TARGET_CLAUSE_BOILERPLATE.sub("", clause)
        remainder = re.sub(r"[\s：:，,、（）()\[\]{}]+", "", remainder)
        if len(remainder) >= 2:
            return True
    return False


def _pleading_reference_has_locator(value: Any) -> bool:
    """Require a page, paragraph, item, or named pleading section."""

    return isinstance(value, str) and bool(PLEADING_LOCATOR.search(value))


def _proof_object_is_material_summary(value: Any) -> bool:
    """Return True when the Word proof object merely names a carrier/material."""

    raw = _display_proof_object(value)
    if not raw:
        return False
    leading = re.split(
        r"[，,]?\s*(?:并用于|并支持|以支持|用于|对应|服务于)",
        raw,
        maxsplit=1,
    )[0]
    first_clause = re.split(r"[；;。！？!?\n]", leading, maxsplit=1)[0]
    compact = re.sub(r"[\s：:，,、（）()\[\]{}]+", "", first_clause)
    return bool(PROOF_OBJECT_MATERIAL_SUMMARY.fullmatch(compact))


def _proof_object_has_target_anchor(value: Any) -> bool:
    """Require a visible request/defense/element/context reason in Word."""

    return _target_clause_has_specific_content(value, display=True)


def _as_list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _intake_mentions_category(value: Any, category: Any) -> bool:
    return (
        isinstance(value, str)
        and isinstance(category, str)
        and category in INTAKE_CATEGORY_PATTERNS
        and bool(re.search(INTAKE_CATEGORY_PATTERNS[category], value))
    )


def _intake_choice_marker(choice: str) -> str:
    return {
        "unknown": "用户选择不知道",
        "declined": "用户选择不提供",
        "later": "用户选择稍后补充",
    }.get(choice, "")


def _intake_registered_source_ids(data: dict[str, Any]) -> set[str]:
    """Collect material identifiers that the current Manifest actually registers."""

    identifiers: set[str] = set()
    for context in _as_list(data.get("context_materials")):
        if isinstance(context, dict) and _nonempty(context.get("material_id")):
            identifiers.add(_text(context.get("material_id")).strip())
    for group in _as_list(data.get("evidence_groups")):
        if not isinstance(group, dict):
            continue
        for material in _as_list(group.get("materials")):
            if isinstance(material, dict) and _nonempty(material.get("material_id")):
                identifiers.add(_text(material.get("material_id")).strip())
    for source in _as_list(data.get("media_sources")):
        if isinstance(source, dict) and _nonempty(source.get("source_id")):
            identifiers.add(_text(source.get("source_id")).strip())
    return identifiers


def _intake_source_id(value: Any) -> str:
    """Return the registered ID prefix from ``M001:p1:region``-style refs."""

    text = _text(value).strip()
    return re.split(r"[:/#@]", text, maxsplit=1)[0].strip() if text else ""


def _binding_value_present(value: Any) -> bool:
    return value is not None and value != "" and value != [] and value != {}


def _intake_question_topics(value: Any) -> set[str]:
    """Return coarse topics only to catch obviously bundled intake questions."""

    text = _text(value)
    patterns = {
        "amount": r"金额|数额|数值|总额|余额|价款",
        "period": r"期间|起止|时间范围|日期区间",
        "payment": r"付款状态|支付状态|返还状态|结算状态|结清情况|欠付情况|清偿情况|付款情况|支付情况|返还情况|已付未付|是否付款|是否支付|是否返还|是否收回",
        "signer": r"签署人|签字人|谁签署|谁签字",
        "carrier": r"原始载体|载体位置|原件位置|文件位置|手机位置|应用位置",
    }
    return {name for name, pattern in patterns.items() if re.search(pattern, text)}


def _intake_question_bundles_same_topic(value: Any) -> bool:
    """Catch two requested fields hidden inside one broad intake category."""

    text = _text(value)
    patterns = (
        r"(?:金额|数额|主张额).{0,20}(?:以及|并且|同时|和|及|、).{0,20}(?:口径|含义|总额|余额|已付)",
        r"(?:总额|余额|已付金额).{0,20}(?:以及|并且|同时|和|及|、).{0,20}(?:总额|余额|已付金额)",
        r"(?:付款|支付|返还)(?:状态)?.{0,20}(?:以及|并且|同时|和|及|、).{0,20}(?:期间|时间范围|截至时间)",
    )
    return any(re.search(pattern, text) for pattern in patterns)


def _safe_arithmetic_result(value: Any) -> float | None:
    """Evaluate a tiny decimal +/- expression without using eval."""

    if not isinstance(value, str):
        return None
    compact = value.replace(",", "").replace("，", "").replace(" ", "")
    if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?(?:[-+]\d+(?:\.\d+)?)*", compact):
        return None
    return sum(float(token) for token in re.findall(r"[-+]?\d+(?:\.\d+)?", compact))


def _raw_choice_matches(value: Any, choice: str) -> bool:
    """Require only minimal semantic agreement for non-provided choices."""

    text = _text(value)
    if choice == "material_review" and any(
        re.search(pattern, text)
        for pattern in (
            r"(?:不要|不让|无需|不用|别|拒绝).{0,16}(?:由|让|请)?(?:AI|人工智能|模型|你)"
            r".{0,24}(?:材料|附件).{0,16}(?:判断|判定|决定|确定|核对|分析|查|看)",
            r"(?:AI|人工智能|模型|你).{0,12}(?:不能|不要|不应|无需|不用|无法|别)"
            r".{0,16}(?:根据|依据|通过|按|回查|查阅|查看|读取|查)?.{0,8}"
            r"(?:材料|附件).{0,16}(?:判断|判定|决定|确定|核对|分析|查|看)",
        )
    ):
        return False
    patterns = {
        "material_review": (
            r"(?:请|让|由)?(?:AI|人工智能|模型|你).{0,16}(?:自己|自行)?.{0,8}"
            r"(?:回查|查阅|查看|读取|查|根据|依据|通过|按).{0,12}"
            r"(?:现有|已授权|已有|上述|这些|相关)?(?:材料|附件)"
            r"(?:.{0,16}(?:判断|判定|决定|确定|核对|分析))?|"
            r"(?:根据|依据|通过|按|回查|查阅|查看|读取|查).{0,12}"
            r"(?:现有|已授权|已有|上述|这些|相关)?(?:材料|附件).{0,16}"
            r"(?:由|让|请)?(?:AI|人工智能|模型|你)?.{0,8}(?:判断|判定|决定|确定|核对|分析)"
        ),
        "unknown": r"不知道|不清楚|无法确认|不能确认|无法确定|不能确定|尚不确定",
        "declined": r"不提供|不愿提供|拒绝提供|不便提供|不说|不披露",
        "later": r"稍后|以后|后续|之后提供|待补|暂缓|先空着|先不填",
    }
    pattern = patterns.get(choice)
    return pattern is None or bool(re.search(pattern, text))


def _requests_internal_review_draft(value: Any) -> bool:
    return bool(
        re.search(
            r"(?:先|先按|先做|先出|先形成|暂先).{0,12}(?:内部核对稿|核对稿|草稿)|"
            r"(?:内部核对稿|核对稿).{0,8}(?:先做|先出|制作|形成)",
            _text(value),
        )
    )


def _requests_continue_build(value: Any) -> bool:
    """Require an affirmative, reviewable build instruction for stop closure."""

    text = _text(value).strip()
    if not text or _requests_internal_review_draft(text):
        return False
    if re.search(
        r"(?:不要|不必|不用|暂不|先不|停止|取消|不再).{0,12}(?:制作|生成|整理|出稿|输出|执行|继续)|"
        r"(?:制作|生成|整理|出稿|输出).{0,6}(?:停止|取消)",
        text,
    ):
        return False
    return bool(
        re.search(
            r"(?:请|可以|直接|现在|立即|开始|继续|就|按.{0,12})(?:开始|继续|直接)?"
            r"(?:制作|生成|整理|出稿|输出|执行)|"
            r"(?:制作|生成|整理|出稿|输出|执行)(?:吧|即可|可以|继续|开始)|"
            r"(?:开始|继续)(?:做|办理|处理)",
            text,
        )
    )


def _intake_locator_errors(value: Any, label: str) -> list[str]:
    """Validate a review locator without inventing a host message id."""

    if not isinstance(value, dict):
        return [f"{label} 必须是对象"]
    errors: list[str] = []
    kind = _text(value.get("kind")).strip()
    reference = _text(value.get("reference")).strip()
    if not kind:
        errors.append(f"{label}.kind 必须是非空字符串")
    elif kind not in INTAKE_LOCATOR_KINDS:
        errors.append(f"{label}.kind 无效：{kind!r}")
    if not reference:
        errors.append(f"{label}.reference 必须是可回查位置")
    if "message_id" not in value:
        errors.append(f"{label} 必须显式保存 message_id；宿主未提供时写 null")
    elif value.get("message_id") is not None and (
        not isinstance(value.get("message_id"), str) or not value.get("message_id", "").strip()
    ):
        errors.append(f"{label}.message_id 只能是非空字符串或 null；宿主未提供时必须为 null")
    if reference and value.get("message_id") is None:
        transcript_ref = bool(re.fullmatch(r"[^/]+/turn-\d+/(?:user|assistant)", reference))
        conversation_ref = bool(
            re.search(r"第\s*\d+\s*轮.{0,12}(?:用户|助手|assistant|user)", reference, re.IGNORECASE)
            or re.fullmatch(r"[^/]+/[^/]+(?:/[^/]+)*", reference)
        )
        if kind == "transcript_turn" and not transcript_ref:
            errors.append(f"{label}.reference 必须使用 scenario/turn-N/user|assistant 格式")
        if kind == "conversation_turn" and not conversation_ref:
            errors.append(f"{label}.reference 必须明确轮次/角色或使用可解释的分段定位")
    return errors


def _intake_locator_key(value: Any) -> tuple[str, str, str | None] | None:
    if not isinstance(value, dict):
        return None
    kind = _text(value.get("kind")).strip()
    reference = _text(value.get("reference")).strip()
    if not kind or not reference:
        return None
    message_id = value.get("message_id") if isinstance(value.get("message_id"), str) else None
    return kind, reference, message_id


def _field_binding_amount_errors(item: dict[str, Any]) -> list[str]:
    """Validate only explicit amount arithmetic in a bound answer.

    Intake does not verify facts.  It may, however, preserve a user''s selected
    amount basis and reject an explicit second deduction of a payment already
    included in ``net_outstanding``.
    """

    field_key = _text(item.get("field_key")).lower()
    is_payment_field = any(token in field_key for token in ("payment", "refund", "付款", "支付", "返还", "清偿"))
    is_amount_field = any(token in field_key for token in ("amount", "金额", "款项", "gross_due", "net_outstanding"))
    if not (is_amount_field or is_payment_field):
        return []
    binding = item.get("answer_binding")
    if not isinstance(binding, dict):
        return []
    value = binding.get("normalized_value")
    errors: list[str] = []
    if is_payment_field:
        if binding.get("resolution_status") == "resolved":
            if not isinstance(value, dict) or not _nonempty(value.get("time_range") or value.get("payment_time_range")):
                errors.append("付款状态必须绑定具体 entity_ref 及时间范围")
            elif _nonempty(value.get("entity_ref")) and _text(value.get("entity_ref")).strip() != _text(item.get("entity_ref")).strip():
                errors.append("付款状态 normalized_value.entity_ref 必须与问题 entity_ref 一致")
    if not isinstance(value, dict):
        return errors
    numeric_keys = (
        "source_amount",
        "line_item_amount",
        "computed_amount",
        "adopted_amount",
        "gross_due",
        "net_outstanding",
        "historical_paid",
        "historical_payment_amount",
    )
    for key in numeric_keys:
        if key not in value:
            continue
        numeric = _numeric_amount(value.get(key))
        if numeric is None or numeric < 0:
            errors.append(f"field_binding_v1 的金额字段 {key} 必须是非负数")
    amount_kind = _text(value.get("amount_kind") or value.get("basis")).strip()
    if amount_kind in {"net_outstanding", "balance_after_historical_payment"} and value.get("historical_payments_already_deducted") is not True:
        errors.append("net_outstanding 必须明确 historical_payments_already_deducted=true")
    if value.get("historical_payments_already_deducted") is True:
        if value.get("deduct_historical_payments") is not False:
            errors.append("net_outstanding 已标明历史付款已扣除，必须明确 deduct_historical_payments=false")
    amount_selection = (
        field_key in {"amount.claimed", "amount.adopted", "claimed_amount", "adopted_amount"}
        or "adopted_amount" in value
    )
    if amount_selection and binding.get("resolution_status") == "resolved":
        source_amounts = value.get("source_amounts")
        calculations = value.get("calculation_results")
        if not isinstance(source_amounts, list) or not source_amounts:
            errors.append("采用金额必须保留 source_amounts 及各自来源")
        elif any(
            not isinstance(source, dict)
            or _numeric_amount(source.get("amount")) is None
            or not _nonempty(source.get("source_ref"))
            for source in source_amounts
        ):
            errors.append("source_amounts 每项必须包含 amount 与 source_ref")
        if not isinstance(calculations, list):
            errors.append("采用金额必须保留 calculation_results 数组；没有分项计算时使用空数组")
        else:
            calculation_values: list[float] = []
            for index, calculation in enumerate(calculations, start=1):
                if not isinstance(calculation, dict):
                    errors.append(f"calculation_results 第{index}项必须是对象")
                    continue
                computed = _safe_arithmetic_result(calculation.get("expression"))
                declared = _numeric_amount(calculation.get("result"))
                if computed is None or declared is None:
                    errors.append(f"calculation_results 第{index}项必须保存可复算的加减表达式和数值结果")
                    continue
                source_refs = calculation.get("source_refs")
                if not isinstance(source_refs, list) or not source_refs or any(not _nonempty(ref) for ref in source_refs):
                    errors.append(f"calculation_results 第{index}项必须保存非空 source_refs")
                if abs(computed - declared) > 0.005:
                    errors.append(f"calculation_results 第{index}项的 expression 与 result 不一致")
                calculation_values.append(declared)
            selected_net = _numeric_amount(value.get("net_outstanding"))
            if amount_kind == "net_outstanding":
                if not calculation_values:
                    errors.append("net_outstanding 必须保留至少一项可复算的 calculation_results")
                else:
                    if selected_net is None:
                        selected_net = _numeric_amount(value.get("adopted_amount"))
                    if selected_net is not None and any(
                        abs(result - selected_net) > 0.005 for result in calculation_values
                    ):
                        errors.append(
                            "resolved net_outstanding 的所有 calculation_results 必须与采用金额一致；"
                            "差异计算应保留来源并将该字段标为 ambiguous"
                        )
        if _numeric_amount(value.get("adopted_amount")) is None:
            errors.append("采用金额必须单独记录 adopted_amount")
        if amount_kind not in {"gross_due", "net_outstanding"}:
            errors.append("采用金额的 amount_kind 必须是 gross_due 或 net_outstanding")
    deductions = value.get("deduction_operations")
    if isinstance(deductions, list):
        seen_deductions: set[str] = set()
        for deduction in deductions:
            if not isinstance(deduction, dict):
                continue
            deduction_id = _text(deduction.get("payment_id") or deduction.get("entity_ref")).strip()
            if deduction_id and deduction_id in seen_deductions:
                errors.append(f"历史付款 {deduction_id} 在同一金额绑定中被重复扣减")
            if deduction_id:
                seen_deductions.add(deduction_id)
    gross_due = _numeric_amount(value.get("gross_due"))
    net_outstanding = _numeric_amount(value.get("net_outstanding"))
    historical_paid = _numeric_amount(value.get("historical_paid", value.get("historical_payment_amount")))
    if gross_due is not None and net_outstanding is not None and historical_paid is not None:
        expected = gross_due - historical_paid
        if abs(expected - net_outstanding) > 0.005:
            errors.append("gross_due、historical_paid 与 net_outstanding 的分项计算结果不一致")
    return errors


def _field_binding_question_errors(
    item: dict[str, Any],
    *,
    category: str,
    choice: str,
    seen_question_ids: set[str],
    binding_statuses: dict[str, str],
    field_history: dict[tuple[str, str], str],
    registered_source_ids: set[str],
) -> tuple[list[str], str]:
    """Validate one field_binding_v1 question without inventing an answer."""

    errors: list[str] = []
    label = "field_binding_v1 问题"
    question_id = item.get("question_id")
    if not isinstance(question_id, str) or not question_id.strip():
        errors.append(f"{label}必须有非空 question_id")
        question_id = ""
    else:
        question_id = question_id.strip()
        if question_id in seen_question_ids:
            errors.append(f"field_binding_v1 question_id 重复：{question_id}")
    entity_ref = item.get("entity_ref")
    if not isinstance(entity_ref, str) or not entity_ref.strip():
        errors.append(f"{label}必须明确 entity_ref")
        entity_ref = ""
    else:
        entity_ref = entity_ref.strip()
    field_key = item.get("field_key")
    if not isinstance(field_key, str) or not field_key.strip():
        errors.append(f"{label}必须明确 field_key")
        field_key = ""
    else:
        field_key = field_key.strip()
    question = item.get("question")
    if not isinstance(question, str) or not question.strip():
        errors.append(f"{label}缺少实际提问文本")
    else:
        if "\n" in question or "\r" in question or ";" in question or "；" in question:
            errors.append(f"{label}必须一次只处理一个明确事项")
        if len(re.findall(r"[?？]", question)) > 1:
            errors.append(f"{label}不得包含多个独立问题")
        if len(_intake_question_topics(question)) >= 2 or _intake_question_bundles_same_topic(question):
            errors.append(f"{label}把金额、期间、付款、签署人或载体等多个事项压在同一问题中")
    if category not in INTAKE_CATEGORIES:
        errors.append(f"{label}的 category 无效：{category!r}")
    elif not _intake_mentions_category(question, category):
        errors.append(f"{label}的提问文本未对应 category={category}")

    depends_on = item.get("depends_on", [])
    if not isinstance(depends_on, list):
        errors.append(f"{label}.depends_on 必须是 question_id 数组")
        depends_on = []
    if len(depends_on) != len(set(depends_on)):
        errors.append(f"{label}.depends_on 不得重复")
    for dependency in depends_on:
        if not isinstance(dependency, str) or not dependency.strip():
            errors.append(f"{label}.depends_on 只能包含非空 question_id")
            continue
        dependency = dependency.strip()
        if dependency not in seen_question_ids:
            errors.append(f"{label}的依赖问题必须先出现：{dependency}")
        elif binding_statuses.get(dependency) != "resolved":
            errors.append(f"{label}的依赖问题尚未 resolved：{dependency}")

    context = item.get("answer_context")
    context_kind = ""
    context_version = ""
    context_options: list[Any] = []
    option_ids: list[str] = []
    option_labels: list[str] = []
    option_values: dict[str, Any] = {}
    summary_fields: list[str] = []
    summary_values: dict[str, Any] = {}
    context_summary_id = ""
    binary_proposition = ""
    summary_context_locator: Any = None
    if not isinstance(context, dict):
        errors.append(f"{label}缺少 answer_context 对象")
    else:
        context_kind = _text(context.get("kind")).strip()
        if context_kind not in ANSWER_CONTEXT_KINDS:
            errors.append(f"{label}.answer_context.kind 无效：{context_kind!r}")
        version = context.get("version")
        if not isinstance(version, str) or not version.strip():
            errors.append(f"{label}.answer_context.version 必须是非空版本字符串")
        else:
            context_version = version.strip()
        raw_summary_id = context.get("summary_id")
        if raw_summary_id is not None:
            context_summary_id = _text(raw_summary_id).strip()
        raw_proposition = context.get("proposition")
        if raw_proposition is not None:
            binary_proposition = _text(raw_proposition).strip()
        options = context.get("options")
        if not isinstance(options, list):
            errors.append(f"{label}.answer_context.options 必须是数组")
            options = []
        else:
            for option in options:
                if not isinstance(option, dict):
                    errors.append(f"{label}.answer_context.options 每项必须是对象")
                    continue
                option_id = _text(option.get("option_id")).strip()
                option_label = _text(option.get("label")).strip()
                if not option_id or not option_label:
                    errors.append(f"{label}.answer_context.options 每项必须有 option_id 与 label")
                    continue
                if "normalized_value" not in option:
                    errors.append(f"{label}.answer_context.options 每项必须保存 normalized_value")
                    continue
                option_ids.append(option_id)
                option_labels.append(option_label)
                option_values[option_id] = option.get("normalized_value")
            if len(option_ids) != len(set(option_ids)) or len(option_labels) != len(set(option_labels)):
                errors.append(f"{label}.answer_context.options 的 option_id 与 label 不得重复")
        context_options = options
        raw_summary_fields = context.get("summary_fields")
        if not isinstance(raw_summary_fields, list) or any(
            not isinstance(field, str) or not field.strip() for field in raw_summary_fields
        ):
            errors.append(f"{label}.answer_context.summary_fields 必须是字符串数组")
            raw_summary_fields = []
        summary_fields = [field.strip() for field in raw_summary_fields]
        if len(summary_fields) != len(set(summary_fields)):
            errors.append(f"{label}.answer_context.summary_fields 不得重复")
        raw_summary_values = context.get("summary_values")
        if isinstance(raw_summary_values, dict):
            summary_values = raw_summary_values
        if context_kind in {"binary", "options"} and not context_options:
            errors.append(f"{label}的 binary/options answer_context 必须提供 options")
        if context_kind == "binary":
            if len(option_ids) != 2:
                errors.append(f"{label}的 binary answer_context 必须提供两个稳定选项")
            if not binary_proposition:
                errors.append(f"{label}的 binary answer_context 必须保存具体 proposition")
        if context_kind == "summary":
            if not summary_fields:
                errors.append(f"{label}的 summary answer_context 必须提供 summary_fields")
            if not context_summary_id:
                errors.append(f"{label}的 summary answer_context 必须保存 summary_id")
            if not isinstance(raw_summary_values, dict) or set(raw_summary_values) != set(summary_fields):
                errors.append(f"{label}的 summary answer_context 必须保存与 summary_fields 完全对应的 summary_values")
            raw_summary_snapshot = context.get("summary_snapshot")
            if not isinstance(raw_summary_snapshot, str) or not raw_summary_snapshot.strip():
                errors.append(f"{label}的 summary answer_context 必须保存实际展示的 summary_snapshot")
            elif context_summary_id and context_summary_id not in raw_summary_snapshot:
                errors.append(f"{label}的 summary_snapshot 必须包含当前 summary_id")
            summary_context_locator = context.get("context_locator")
            errors.extend(
                _intake_locator_errors(
                    summary_context_locator,
                    f"{label}.answer_context.context_locator",
                )
            )

    binding = item.get("answer_binding")
    if not isinstance(binding, dict):
        errors.append(f"{label}缺少 answer_binding 对象")
        return errors, "unresolved"
    raw_response = binding.get("raw_response")
    if not isinstance(raw_response, str) or not raw_response.strip():
        errors.append(f"{label}.answer_binding.raw_response 必须保留原始答复")
        raw_response = ""
    locator = binding.get("response_locator")
    errors.extend(_intake_locator_errors(locator, f"{label}.answer_binding.response_locator"))
    if context_kind == "summary" and (
        _intake_locator_key(summary_context_locator) is not None
        and _intake_locator_key(summary_context_locator) == _intake_locator_key(locator)
    ):
        errors.append("摘要展示位置不得与用户确认答复位置相同")

    normalized_present = "normalized_value" in binding
    normalized_value = binding.get("normalized_value")
    if not normalized_present:
        errors.append(f"{label}.answer_binding 必须保留 normalized_value 字段")
    binding_basis = binding.get("binding_basis")
    if not isinstance(binding_basis, dict):
        errors.append(f"{label}.answer_binding.binding_basis 必须是对象")
        binding_basis = {}
    basis_kind = _text(binding_basis.get("kind")).strip()
    if basis_kind not in BINDING_BASIS_KINDS:
        errors.append(f"{label}.binding_basis.kind 无效：{basis_kind!r}")
    if binding_basis.get("question_id") != question_id:
        errors.append(f"{label}.binding_basis.question_id 必须回指当前 question_id")
    if binding_basis.get("context_version") != context_version:
        errors.append(f"{label}.binding_basis.context_version 必须对应 answer_context.version")
    raw_response_fragment = binding_basis.get("raw_response_fragment")
    if raw_response_fragment is not None and (
        not isinstance(raw_response_fragment, str)
        or not raw_response_fragment.strip()
        or raw_response_fragment.strip() not in raw_response
    ):
        errors.append("binding_basis.raw_response_fragment 必须是 raw_response 中的非空原文片段")
    if choice == "material_review" and basis_kind not in {"material_lookup", "computed"}:
        errors.append("material_review 必须使用 material_lookup 或 computed 绑定材料判断依据")
    if basis_kind == "material_lookup":
        source_refs = binding_basis.get("source_refs", binding.get("source_refs"))
        if not isinstance(source_refs, list) or not source_refs or any(not _nonempty(ref) for ref in source_refs):
            errors.append(f"材料回查 binding_basis=material_lookup 时必须提供 source_refs")
    if basis_kind == "selected_option":
        selected_option_id = _text(binding_basis.get("selected_option_id")).strip()
        if context_kind not in {"binary", "options"}:
            errors.append("selected_option 只能用于 binary/options answer_context")
        elif not selected_option_id or selected_option_id not in option_ids:
            errors.append("selected_option_id 必须回指当前版本的具体选项")
        elif normalized_value != option_values.get(selected_option_id):
            errors.append("normalized_value 必须与 selected_option_id 对应选项的 normalized_value 完全一致")
    if basis_kind == "computed":
        source_refs = binding_basis.get("source_refs", binding.get("source_refs"))
        if not isinstance(source_refs, list) or not source_refs or any(not _nonempty(ref) for ref in source_refs):
            errors.append("computed 绑定必须提供可回查 source_refs")
    if choice == "material_review":
        source_refs = binding_basis.get("source_refs", binding.get("source_refs"))
        if not registered_source_ids:
            errors.append("material_review 必须引用当前 Manifest 已登记的授权材料")
        elif isinstance(source_refs, list) and source_refs:
            unknown_sources = sorted(
                {
                    _text(source_ref).strip()
                    for source_ref in source_refs
                    if _intake_source_id(source_ref) not in registered_source_ids
                }
            )
            if unknown_sources:
                errors.append(
                    "material_review 的 source_refs 未在当前 Manifest 登记："
                    + "、".join(unknown_sources)
                )
    if basis_kind == "direct_answer" and context_kind in {"binary", "options"}:
        errors.append("binary/options answer_context 必须使用 selected_option 绑定")

    resolution_status = _text(binding.get("resolution_status")).strip()
    if resolution_status not in BINDING_RESOLUTION_STATUSES:
        errors.append(f"{label}.resolution_status 无效：{resolution_status!r}")
    if choice in {"provided", "material_review"}:
        if resolution_status not in {"resolved", "ambiguous"}:
            errors.append(f"{choice} 回答只能使用 resolution_status=resolved 或 ambiguous")
        if resolution_status == "resolved" and not _binding_value_present(normalized_value):
            errors.append(f"{choice} 且 resolved 的回答必须保存非空 normalized_value")
        if resolution_status == "ambiguous" and not _as_list(binding.get("ambiguities")):
            errors.append(f"{choice} 的 ambiguous 回答必须保留 ambiguities")
        if resolution_status == "resolved" and _as_list(binding.get("ambiguities")):
            errors.append(f"{choice} 的 resolved 回答不得同时保留未解决 ambiguities")
    elif resolution_status != choice:
        errors.append(f"{choice} 回答必须保留 resolution_status={choice}")
    if choice not in {"provided", "material_review"} and _binding_value_present(normalized_value):
        errors.append(f"{choice} 回答不得伪造已规范化的事实值")
    if choice != "provided" and not _raw_choice_matches(raw_response, choice):
        errors.append(f"raw_response 与 user_choice={choice} 的语义不一致；应标为 ambiguous 或更正状态")

    ambiguities = binding.get("ambiguities")
    if not isinstance(ambiguities, list):
        errors.append(f"{label}.ambiguities 必须是数组")
    if binding.get("fact_verification") != FACT_VERIFICATION_NOT_VERIFIED:
        errors.append(f"{label}.fact_verification 必须明确为 {FACT_VERIFICATION_NOT_VERIFIED}")

    compact_response = re.sub(r"[\s。；，：:、,.，;！？!?（）()\[\]{}]", "", raw_response)
    short_answers = {
        "是", "否", "好", "好的", "可以", "行", "没问题", "确认", "同意", "对", "正确",
        "第一项", "第二项", "第三项", "按你说的", "按上述",
    }
    if compact_response in short_answers:
        if context_kind not in {"binary", "options", "summary"}:
            errors.append("是/第一项/按你说的等短答只有在明确 binary、options 或 summary 版本中才可绑定")
        elif compact_response in {"按你说的", "按上述"} and context_kind != "summary":
            errors.append("按你说的等短答必须回指明确的 summary")
        elif compact_response in {"第一项", "第二项", "第三项"}:
            ordinal = {"第一项": 0, "第二项": 1, "第三项": 2}[compact_response]
            selected_option_id = _text(binding_basis.get("selected_option_id")).strip()
            if context_kind != "options" or ordinal >= len(option_ids) or selected_option_id != option_ids[ordinal]:
                errors.append(f"短答 {compact_response} 未准确回指当前版本的选项")
        elif context_kind == "binary":
            selected_option_id = _text(binding_basis.get("selected_option_id")).strip()
            if basis_kind != "selected_option" or compact_response not in option_labels:
                errors.append(f"短答 {compact_response} 未回指当前 binary 命题的具体选项")
            elif option_labels.index(compact_response) >= len(option_ids) or selected_option_id != option_ids[option_labels.index(compact_response)]:
                errors.append(f"短答 {compact_response} 未准确回指当前 binary 选项")
        elif context_kind == "summary" and basis_kind != "confirmed_summary":
            errors.append("摘要后的短答必须使用 confirmed_summary 绑定")

    if basis_kind == "confirmed_summary":
        if context_kind != "summary":
            errors.append("confirmed_summary 只能用于 summary answer_context")
        elif binding_basis.get("summary_id") != context_summary_id:
            errors.append("confirmed_summary.summary_id 必须与当前 answer_context.summary_id 完全一致")
        elif not isinstance(normalized_value, dict) or not set(normalized_value).issubset(set(summary_fields)):
            errors.append("confirmed_summary 只能确认 answer_context.summary_fields")
        elif any(normalized_value.get(field) != summary_values.get(field) for field in normalized_value):
            errors.append("confirmed_summary.normalized_value 必须与当前 summary_values 对应值一致")
        elif compact_response in {
            "是", "好", "好的", "可以", "行", "没问题", "确认", "同意", "对", "正确", "按你说的", "按上述",
        } and normalized_value != summary_values:
            errors.append("摘要后的肯定短答必须完整绑定当前版本的 summary_values")
    errors.extend(_field_binding_amount_errors(item))

    if entity_ref and field_key:
        key = (entity_ref, field_key)
        previous_id = field_history.get(key)
        if previous_id:
            supersedes = item.get("supersedes_question_id")
            if supersedes != previous_id:
                errors.append(
                    f"同一 entity_ref+field_key 重复提问必须 supersedes_question_id={previous_id}"
                )
    return errors, resolution_status


def _intake_reference_valid(value: Any, choice: str, category: str | None = None) -> bool:
    """Require reviewable user-answer text instead of an agent assertion."""

    if not isinstance(value, str):
        return False
    compact = re.sub(r"[\s。；，：:、,.，;！？!?（）()\[\]{}]+", "", value)
    if len(compact) < 10 or not re.search(r"用户|委托人|当事人|本任务|本对话", value):
        return False
    if compact in {"用户已确认", "用户明确确认", "用户在本任务中确认", "用户在本对话中确认"}:
        return False
    if category is not None and not _intake_mentions_category(value, category):
        return False
    if choice == "provided":
        category_clauses = [
            clause
            for clause in re.split(r"[。；;，,：:\n]", value)
            if _intake_mentions_category(clause, category)
        ]
        category_pattern = INTAKE_CATEGORY_PATTERNS[category]
        missing_supply = re.compile(
            r"(?:尚未|仍(?:然)?(?:尚未)?|未|没有|暂无|无法|不能|不愿|拒绝)"
            r".{0,8}(?:提供|提交|递交|上传|发送|附上|补充|给出|收到|取得|出示|找到|获得|交|给|传|发|补|收)|"
            r"没(?:有)?.{0,4}(?:提供|提交|递交|上传|发送|附上|补充|给出|收到|取得|出示|找到|获得|交|给|传|发|补|收)|"
            r"(?:缺失|缺少|欠缺|尚缺|未见)"
        )
        conflicting_choice = re.compile(
            r"不知道|不清楚|无法确认|不能确认|尚未核验|未核验|待核|"
            r"不提供|拒绝提供|选择不提供|"
            r"稍后|以后|之后提供|待补|后续补充"
        )
        if not category_clauses or any(
            missing_supply.search(clause) or conflicting_choice.search(clause)
            for clause in category_clauses
        ):
            return False
        positive = r"(?:已(?:经)?(?:提供|上传|补充|提交|发送|附上)|见附件|随附)"
        unambiguous_gap = (
            r"(?:(?:本案|该|完整|相关|具体|电子版|纸质版|合成|补充|"
            r"相关材料|附件|原件|复印件)){0,2}"
        )
        linked_positive = re.compile(
            rf"(?:{positive}){unambiguous_gap}(?:{category_pattern})|"
            rf"(?:{category_pattern}){unambiguous_gap}(?:{positive})"
        )
        return any(linked_positive.search(clause) for clause in category_clauses)
    anchors = {
        "unknown": r"不知道|不清楚|无法确认|不能确认",
        "declined": r"不提供|拒绝提供|选择不提供",
        "later": r"稍后|以后|之后提供|待补|后续补充",
        "waived": r"跳过|无需询问|不必询问|直接整理|直接生成",
    }
    anchor = anchors.get(choice)
    return bool(anchor and re.search(anchor, value))


def _intake_clarification_gate(data: dict[str, Any]) -> tuple[list[str], list[str]]:
    """Block every DOCX until critical missing-information questions are answered."""

    fatal: list[str] = []
    blockers: list[str] = []
    intake = data.get("intake_clarification")
    if not isinstance(intake, dict):
        return (["CLARIFICATION_REQUIRED：缺少 intake_clarification 询问留痕；即使使用 --allow-draft 也不得生成 DOCX"], [])

    status = intake.get("status")
    if status not in INTAKE_STATUSES:
        fatal.append(f"intake_clarification.status 无效：{status!r}")
        return fatal, blockers
    protocol = intake.get("protocol")
    if protocol is not None and protocol not in INTAKE_PROTOCOLS:
        fatal.append(f"intake_clarification.protocol 无效：{protocol!r}")
        return fatal, blockers

    required = intake.get("required_categories")
    if not isinstance(required, list):
        fatal.append("intake_clarification.required_categories 必须是数组")
        required = []
    elif any(not isinstance(item, str) or item not in INTAKE_CATEGORIES for item in required):
        fatal.append("intake_clarification.required_categories 含无效类别")
    elif len(required) != len(set(required)):
        fatal.append("intake_clarification.required_categories 不得重复")

    if status == "pending":
        fatal.append("CLARIFICATION_REQUIRED：关键缺失信息尚未取得用户回答；禁止生成正式版或草稿 DOCX")
        return fatal, blockers

    if status == "not_required":
        if required:
            fatal.append("intake_clarification.status=not_required 时 required_categories 必须为空")
        if not isinstance(intake.get("reason"), str) or not intake.get("reason", "").strip():
            fatal.append("intake_clarification.status=not_required 必须说明 reason")
        validation = data.get("validation", {})
        checks = validation.get("checks", {}) if isinstance(validation, dict) else {}
        unresolved = validation.get("unresolved_issues", []) if isinstance(validation, dict) else []
        claims = data.get("_all_claims", [])
        has_gap = (
            not isinstance(validation, dict)
            or validation.get("status") != "resolved"
            or (isinstance(unresolved, list) and any(_nonempty(item) for item in unresolved))
            or any(isinstance(check, dict) and check.get("status") == "unresolved" for check in checks.values())
            or any(
                isinstance(claim, dict)
                and isinstance(claim.get("purpose_basis"), dict)
                and (
                    claim["purpose_basis"].get("status") == "provisional"
                    or claim["purpose_basis"].get("kind") == "materials_only"
                )
                for claim in claims
            )
        )
        if has_gap:
            fatal.append("CLARIFICATION_REQUIRED：Manifest 仍有关键缺口，不能声明 intake_clarification.status=not_required")
        return fatal, blockers

    if not required:
        fatal.append(f"intake_clarification.status={status} 时 required_categories 不得为空")

    if status == "waived_by_user":
        waiver = intake.get("waiver_reference")
        if not _intake_reference_valid(waiver, "waived"):
            fatal.append("intake_clarification.waiver_reference 必须记录用户明确要求跳过询问的可复核原话或说明")
        blockers.append("用户明确要求跳过缺失信息询问；只能生成待核对草稿")
        data["_forced_draft"] = True
        return fatal, blockers

    rounds = intake.get("rounds")
    if not isinstance(rounds, list) or not rounds:
        return (fatal + ["CLARIFICATION_REQUIRED：status=completed 但没有结构化询问轮次"], blockers)
    asked_categories: set[str] = set()
    seen_round_ids: set[str] = set()
    unresolved_choices: dict[str, set[str]] = {}
    field_binding = protocol == "field_binding_v1"
    seen_question_ids: set[str] = set()
    binding_statuses: dict[str, str] = {}
    field_history: dict[tuple[str, str], str] = {}
    registered_source_ids = _intake_registered_source_ids(data)
    locator_bindings: dict[tuple[str, str, str | None], list[tuple[str, str]]] = {}
    unresolved_bindings: list[tuple[str, str, str, str]] = []
    draft_request_locator: tuple[str, str, str | None] | None = None
    draft_request_raw = ""
    for round_index, round_item in enumerate(rounds, start=1):
        if not isinstance(round_item, dict):
            fatal.append(f"intake_clarification.rounds 第{round_index}轮必须是对象")
            continue
        round_id = round_item.get("round_id")
        if not isinstance(round_id, str) or not round_id.strip() or round_id in seen_round_ids:
            fatal.append(f"intake_clarification.rounds 第{round_index}轮的 round_id 缺失或重复")
        else:
            seen_round_ids.add(round_id)
        questions = round_item.get("questions")
        if not isinstance(questions, list) or not 1 <= len(questions) <= 3:
            fatal.append(f"intake_clarification.rounds 第{round_index}轮必须包含1—3个问题")
            continue
        for question_index, item in enumerate(questions, start=1):
            label = f"intake_clarification 第{round_index}轮第{question_index}问"
            if not isinstance(item, dict):
                fatal.append(f"{label}必须是对象")
                continue
            category = item.get("category")
            choice = item.get("user_choice")
            if category not in INTAKE_CATEGORIES or category not in required:
                fatal.append(f"{label}的 category 无效或未列入 required_categories")
            else:
                asked_categories.add(category)
            if choice not in INTAKE_CHOICES:
                fatal.append(f"{label}的 user_choice 无效：{choice!r}")
            elif field_binding:
                field_errors, resolution_status = _field_binding_question_errors(
                    item,
                    category=category if isinstance(category, str) else "",
                    choice=choice,
                    seen_question_ids=seen_question_ids,
                    binding_statuses=binding_statuses,
                    field_history=field_history,
                    registered_source_ids=registered_source_ids,
                )
                fatal.extend(f"{label}：{error}" for error in field_errors)
                question_id = _text(item.get("question_id")).strip()
                if question_id and question_id not in seen_question_ids:
                    seen_question_ids.add(question_id)
                    binding_statuses[question_id] = resolution_status
                    entity_ref = _text(item.get("entity_ref")).strip()
                    field_key = _text(item.get("field_key")).strip()
                    if entity_ref and field_key:
                        field_history[(entity_ref, field_key)] = question_id
                if resolution_status == "ambiguous":
                    blockers.append(f"{label} 的回答仍有歧义；需先消歧后再生成正式版")
                if resolution_status != "resolved":
                    entity_ref = _text(item.get("entity_ref")).strip()
                    field_key = _text(item.get("field_key")).strip()
                    if entity_ref and field_key:
                        unresolved_bindings.append((_text(category), entity_ref, field_key, resolution_status))
                if choice in {"unknown", "declined", "later"} and isinstance(category, str):
                    unresolved_choices.setdefault(category, set()).add(choice)
                binding = item.get("answer_binding")
                if isinstance(binding, dict):
                    raw_response = _text(binding.get("raw_response"))
                    locator = binding.get("response_locator")
                    if isinstance(locator, dict):
                        locator_key = (
                            _text(locator.get("kind")),
                            _text(locator.get("reference")),
                            locator.get("message_id") if isinstance(locator.get("message_id"), str) else None,
                        )
                        if draft_request_locator is not None and locator_key != draft_request_locator:
                            fatal.append(
                                f"{label}出现在用户明确要求内部核对稿之后；不得继续追加可降级事实追问"
                            )
                        if _requests_internal_review_draft(raw_response) and draft_request_locator is None:
                            draft_request_locator = locator_key
                            draft_request_raw = raw_response
                        review_key = _intake_locator_key(locator)
                        basis = binding.get("binding_basis")
                        fragment = (
                            _text(basis.get("raw_response_fragment")).strip()
                            if isinstance(basis, dict)
                            else ""
                        )
                        if review_key is not None:
                            locator_bindings.setdefault(review_key, []).append((label, fragment))
            elif choice == "material_review":
                fatal.append(f"{label}：material_review 仅支持 protocol=field_binding_v1，旧问答不得自动补造材料判断绑定")
            elif not _intake_reference_valid(item.get("response_reference"), choice, category):
                fatal.append(f"{label}缺少可复核的用户回答留痕")
            elif choice in {"unknown", "declined", "later"} and isinstance(category, str):
                unresolved_choices.setdefault(category, set()).add(choice)
    for records in locator_bindings.values():
        if len(records) > 1 and any(not fragment for _, fragment in records):
            labels = "、".join(label for label, _ in records)
            fatal.append(f"同一答复位置绑定多个字段时，每个绑定必须保存 raw_response_fragment：{labels}")
    missing_coverage = set(required) - asked_categories
    if missing_coverage:
        fatal.append(f"CLARIFICATION_REQUIRED：以下缺失类别尚未完成询问：{sorted(missing_coverage)}")
    unresolved_issues = data.get("validation", {}).get("unresolved_issues", []) if isinstance(data.get("validation"), dict) else []
    issue_categories = {
        match.group(1)
        for item in unresolved_issues if isinstance(unresolved_issues, list)
        for match in re.finditer(r"信息缺失\[([^\]]+)\]", _text(item))
    }
    unasked_issues = issue_categories - asked_categories
    if unasked_issues:
        fatal.append(f"CLARIFICATION_REQUIRED：以下未解决缺失类别没有对应用户问答：{sorted(unasked_issues)}")
    for category, choices in unresolved_choices.items():
        for choice in choices:
            category_marker = f"信息缺失[{category}]"
            choice_marker = _intake_choice_marker(choice)
            if field_binding:
                kept = isinstance(unresolved_issues, list) and any(
                    category_marker in _text(issue)
                    and (choice_marker in _text(issue) or choice in _text(issue).lower())
                    for issue in unresolved_issues
                )
            else:
                kept = isinstance(unresolved_issues, list) and any(
                    category_marker in _text(issue) for issue in unresolved_issues
                )
            if not kept:
                fatal.append(f"用户对[{category}]选择{choice}，但 validation.unresolved_issues 未保留该状态")
    if field_binding:
        unresolved_status_patterns = {
            "ambiguous": r"ambiguous|歧义|待消歧",
            "unknown": r"unknown|不知道|不清楚|无法确认|不能确认",
            "declined": r"declined|不提供|拒绝",
            "later": r"later|稍后|后续|待补|暂缓",
        }
        for category, entity_ref, field_key, resolution_status in unresolved_bindings:
            status_pattern = unresolved_status_patterns.get(resolution_status)
            if status_pattern is None:
                continue
            if not isinstance(unresolved_issues, list) or not any(
                f"信息缺失[{category}]" in _text(issue)
                and entity_ref in _text(issue)
                and field_key in _text(issue)
                and bool(re.search(status_pattern, _text(issue), re.IGNORECASE))
                for issue in unresolved_issues
            ):
                fatal.append(
                    f"field_binding_v1 的未解决字段未完整留痕：category={category}, entity_ref={entity_ref}, "
                    f"field_key={field_key}, status={resolution_status}"
                )
        stop_decision = intake.get("stop_decision")
        if not isinstance(stop_decision, dict):
            fatal.append("field_binding_v1 在 completed 时必须记录 stop_decision")
        else:
            mode = stop_decision.get("mode")
            if mode not in {"continue_build", "internal_review_draft"}:
                fatal.append("field_binding_v1 stop_decision.mode 必须是 continue_build 或 internal_review_draft")
            if not isinstance(stop_decision.get("reason"), str) or not stop_decision.get("reason", "").strip():
                fatal.append("field_binding_v1 stop_decision 必须说明 reason")
            if not isinstance(stop_decision.get("raw_response"), str) or not stop_decision.get("raw_response", "").strip():
                fatal.append("field_binding_v1 stop_decision 必须保留触发停止或制作的用户原话")
                stop_raw_response = ""
            else:
                stop_raw_response = stop_decision["raw_response"].strip()
                if mode == "continue_build" and not _requests_continue_build(stop_raw_response):
                    fatal.append("field_binding_v1 continue_build 必须回指用户明确要求制作、生成或继续执行的原话")
                if mode == "internal_review_draft" and not _requests_internal_review_draft(stop_raw_response):
                    fatal.append("field_binding_v1 internal_review_draft 必须回指用户明确要求内部核对稿或草稿的原话")
            locator = stop_decision.get("response_locator")
            fatal.extend(
                _intake_locator_errors(
                    locator,
                    "field_binding_v1 stop_decision.response_locator",
                )
            )
            unresolved_fields = stop_decision.get("unresolved_fields")
            if not isinstance(unresolved_fields, list) or any(
                not isinstance(field, str) or not field.strip() for field in unresolved_fields
            ):
                fatal.append("field_binding_v1 stop_decision.unresolved_fields 必须是非空字符串组成的数组")
                unresolved_fields = []
            required_unresolved = {
                f"{entity_ref}:{field_key}"
                for _, entity_ref, field_key, _ in unresolved_bindings
            }
            if not required_unresolved.issubset(set(unresolved_fields)):
                fatal.append("field_binding_v1 stop_decision 未列全仍待核的对象字段")
            if unresolved_bindings:
                blockers.append("field_binding_v1 仍有歧义、未知、拒绝或暂缓字段；不得生成正式版")
            if mode == "internal_review_draft":
                blockers.append("用户明确要求先做内部核对稿；保留待核事项并强制草稿")
                data["_forced_draft"] = True
                data["_intake_draft_authorized"] = True
            if draft_request_locator is not None:
                stop_locator = stop_decision.get("response_locator")
                stop_locator_key = (
                    _text(stop_locator.get("kind")) if isinstance(stop_locator, dict) else "",
                    _text(stop_locator.get("reference")) if isinstance(stop_locator, dict) else "",
                    stop_locator.get("message_id") if isinstance(stop_locator, dict) and isinstance(stop_locator.get("message_id"), str) else None,
                )
                if mode != "internal_review_draft":
                    fatal.append("用户原话已明确要求内部核对稿，stop_decision.mode 必须为 internal_review_draft")
                if stop_locator_key != draft_request_locator or stop_decision.get("raw_response") != draft_request_raw:
                    fatal.append("stop_decision 必须回指并原样保存用户提出内部核对稿的同一轮答复")
    return fatal, blockers


def _normalise_check(value: Any) -> dict[str, Any]:
    if isinstance(value, bool):
        return {
            "status": "resolved" if value else "unresolved",
            "material_refs": [],
            "note": "旧版布尔核验字段，缺少逐项材料依据。" if value else "旧版布尔核验字段未确认。",
        }
    if not isinstance(value, dict):
        return {"status": "unresolved", "material_refs": [], "note": "核验字段不是记录对象。"}
    return {
        "status": value.get("status", "unresolved"),
        "material_refs": value.get("material_refs", []),
        "note": _text(value.get("note", "")),
    }


def normalise_manifest(data: dict[str, Any]) -> dict[str, Any]:
    """Make v1 and v2 shapes consumable while preserving all user fields."""

    normalized = copy.deepcopy(data)
    if "schema_version" not in data:
        source_version = 1
    else:
        raw_version = data.get("schema_version")
        if type(raw_version) is int and raw_version == 1:
            source_version = 1
        elif type(raw_version) is int and raw_version == 2:
            source_version = 2
        else:
            raise ManifestError(
                [
                    f"schema_version 类型或取值无效：{raw_version!r}；必须是整数1或2，"
                    "布尔值、浮点数、字符串及显式 null 均不接受"
                ]
            )
    strict_v2 = source_version == 2
    normalized["_source_schema_version"] = source_version
    normalized["_strict_v2"] = strict_v2
    normalized["_forced_draft"] = source_version == 1

    context_materials = normalized.get("context_materials", [])
    if not isinstance(context_materials, list):
        if strict_v2 and "context_materials" in normalized:
            normalized["_invalid_context_materials_type"] = True
        context_materials = []
    normalized["context_materials"] = copy.deepcopy(context_materials)

    groups = normalized.get("evidence_groups")
    if not isinstance(groups, list):
        if strict_v2 and "evidence_groups" in normalized:
            normalized["_invalid_evidence_groups_type"] = True
        groups = []
    material_counter = 1
    claim_counter = 1
    all_claims: list[dict[str, Any]] = []
    normalized_groups: list[dict[str, Any]] = []
    for group_index, raw_group in enumerate(groups, start=1):
        group = raw_group if isinstance(raw_group, dict) else {}
        group = copy.deepcopy(group)
        if not strict_v2 and not _nonempty(group.get("number")):
            group["number"] = group_index
        if not _nonempty(group.get("group_id")):
            # group_id is a technical ownership key, unlike evidence content
            # fields; derive it deterministically when omitted so legacy v2
            # callers can still obtain a one-time claim owner.
            group["group_id"] = f"G{group_index:02d}"
            if strict_v2:
                group["_generated_group_id"] = True
        if not strict_v2 and not _nonempty(group.get("evidence_name")):
            group["evidence_name"] = "（待补充证据名称）"
        if not strict_v2 and not _nonempty(group.get("evidence_form")):
            group["evidence_form"] = "（待核对证据形式）"
        if not strict_v2 and not _nonempty(group.get("proof_object")):
            group["proof_object"] = "（待补充证明对象）"
        materials = group.get("materials")
        if not isinstance(materials, list):
            if strict_v2 and "materials" in group:
                group["_invalid_materials_type"] = True
            materials = []
        normalized_materials: list[dict[str, Any]] = []
        for raw_material in materials:
            material = raw_material if isinstance(raw_material, dict) else {}
            material = copy.deepcopy(material)
            if not _nonempty(material.get("material_id")) and not strict_v2:
                material["material_id"] = f"M{material_counter:03d}"
            material_counter += 1
            normalized_materials.append(material)
        group["materials"] = normalized_materials

        group_claims = group.get("proof_claims", [])
        if group_claims is None or not isinstance(group_claims, list):
            if strict_v2 and "proof_claims" in group:
                group["_invalid_proof_claims_type"] = True
            group_claims = []
        normalized_group_claims: list[dict[str, Any]] = []
        for raw_claim in group_claims:
            claim = raw_claim if isinstance(raw_claim, dict) else raw_claim
            claim = copy.deepcopy(claim)
            if isinstance(claim, dict) and not _nonempty(claim.get("claim_id")) and not strict_v2:
                claim["claim_id"] = f"C{claim_counter:03d}"
            claim_counter += 1
            if isinstance(claim, dict):
                claim["_container_group_id"] = _text(group.get("group_id"))
            normalized_group_claims.append(claim)
            all_claims.append(claim)
        group["proof_claims"] = normalized_group_claims
        normalized_groups.append(group)

    top_claims = normalized.get("proof_claims", [])
    if top_claims is None or not isinstance(top_claims, list):
        if strict_v2 and "proof_claims" in normalized:
            normalized["_invalid_top_proof_claims_type"] = True
        top_claims = []
    normalized_top_claims: list[dict[str, Any]] = []
    for raw_claim in top_claims:
        claim = raw_claim if isinstance(raw_claim, dict) else raw_claim
        claim = copy.deepcopy(claim)
        if isinstance(claim, dict) and not _nonempty(claim.get("claim_id")) and not strict_v2:
            claim["claim_id"] = f"C{claim_counter:03d}"
        claim_counter += 1
        normalized_top_claims.append(claim)
        all_claims.append(claim)
    normalized["proof_claims"] = normalized_top_claims
    normalized["_all_claims"] = all_claims
    normalized["evidence_groups"] = normalized_groups

    validation = normalized.get("validation")
    if not isinstance(validation, dict):
        if strict_v2 and "validation" in normalized:
            normalized["_invalid_validation_type"] = True
        validation = {}
    validation = copy.deepcopy(validation)
    raw_checks = validation.get("checks")
    if not isinstance(raw_checks, dict):
        if strict_v2 and "checks" in validation:
            normalized["_invalid_checks_type"] = True
        raw_checks = {}
    legacy_boolean = False
    checks: dict[str, dict[str, Any]] = {}
    for key in CHECK_KEYS:
        if isinstance(raw_checks.get(key), bool):
            legacy_boolean = True
        checks[key] = _normalise_check(raw_checks.get(key))
    validation["checks"] = checks
    if "unresolved_issues" not in validation:
        validation["unresolved_issues"] = []
    elif not isinstance(validation.get("unresolved_issues"), list):
        normalized["_invalid_unresolved_issues_type"] = True
    normalized["validation"] = validation
    normalized["_legacy_boolean_checks"] = legacy_boolean or source_version == 1
    return normalized


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _display_rotation(value: Any) -> int:
    """Validate the user-declared clockwise display rotation.

    EXIF orientation is handled separately and automatically; this field is
    the only manual correction accepted by the bundle builder.  Do not coerce
    strings or guess a value from image dimensions.
    """

    if value is None:
        return 0
    if isinstance(value, bool) or not isinstance(value, int) or value not in ALLOWED_DISPLAY_ROTATIONS:
        raise ValueError("display_rotation_degrees_clockwise 只接受整数 0、90、180 或 270")
    return value


def _image_preflight(path: Path, declared_readability: Any = None, display_rotation_degrees_clockwise: int = 0) -> dict[str, Any]:
    """Verify bytes, then reopen/load; do not trust extension or metadata alone."""

    try:
        display_rotation_degrees_clockwise = _display_rotation(display_rotation_degrees_clockwise)
    except ValueError as exc:
        return {
            "path": path,
            "decode_status": "error",
            "readability": "unreadable",
            "warnings": [],
            "errors": [str(exc)],
            "exif_orientation": None,
            "exif_transposed": False,
            "display_rotation_degrees_clockwise": display_rotation_degrees_clockwise,
        }
    record: dict[str, Any] = {
        "path": path,
        "decode_status": "error",
        "readability": "unreadable",
        "warnings": [],
        "errors": [],
        "exif_orientation": None,
        "exif_transposed": False,
        "display_rotation_degrees_clockwise": display_rotation_degrees_clockwise,
    }
    try:
        from PIL import Image, ImageOps

        # Read dimensions without loading pixels; extreme images are reported
        # and blocked before any potentially unsafe allocation.
        Image.MAX_IMAGE_PIXELS = None

        # Keep verify immediately after open. Some Pillow wrappers and the
        # acceptance harness deliberately enforce this ordering.
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            record["format"] = image.format
            record["width"], record["height"] = image.size
            record["exif_orientation"] = image.getexif().get(274)
            if max(int(record["width"]), int(record["height"])) > 200_000 or int(record["width"]) * int(record["height"]) > 200_000_000:
                record["errors"].append("图片尺寸极端，拒绝加载以避免内存风险")
                return record
            image.load()
            normalized = ImageOps.exif_transpose(image)
            normalized.load()
            if display_rotation_degrees_clockwise:
                normalized = normalized.rotate(-display_rotation_degrees_clockwise, expand=True)
                normalized.load()
            record["normalized_width"], record["normalized_height"] = normalized.size
            record["mode"] = normalized.mode
        width = int(record["normalized_width"])
        height = int(record["normalized_height"])
        record["decode_status"] = "ok"
        record["exif_transposed"] = bool(record["exif_orientation"] not in (None, 1))
        max_dimension = max(width, height)
        pixels = width * height
        record["pixels"] = pixels
        ratio = max(width, height) / max(1, min(width, height))
        record["aspect_ratio"] = round(ratio, 3)
        if max_dimension > 200_000 or pixels > 200_000_000:
            record["errors"].append("图片尺寸极端，无法安全嵌入")
        elif max_dimension > 100_000 or pixels > 100_000_000:
            record["warnings"].append("图片尺寸很大，可能造成内存或分页风险")
        if ratio >= 8:
            record["warnings"].append("长图：缩放后文字可能不可读，建议全貌页加局部页")
        if min(width, height) < 600 or pixels < 500_000:
            record["warnings"].append("源像素较低，缩放后存在不可读风险")
        declared = _text(declared_readability).lower()
        if declared in {"poor", "unreadable", "不可读", "模糊"}:
            record["errors"].append("材料标注为不可读或严重模糊")
        elif declared in {"partial", "部分可读", "模糊部分"}:
            record["errors"].append("材料仅部分可读；正式版需补充全貌页和局部清晰页")
        record["readability"] = "unreadable" if record["errors"] else ("warning" if record["warnings"] else "clear")
    except ImportError as exc:
        record["errors"].append(f"缺少 Pillow，无法实际解码图片：{exc}")
    except Exception as exc:  # Pillow raises several image-specific exceptions.
        record["errors"].append(f"图片损坏或不可读（corrupt image）：图片解码失败：{exc}")
    return record


def _source_value(material: dict[str, Any]) -> Any:
    """Provenance label (for example a recording or screenshot source)."""

    return material.get("source")


_CARRIER_PLACEHOLDER_RE = re.compile(
    r"^(?:已核对|已核验|已留存|原件已核对|原件已核验|原始载体已核验|原始载体已核对|见原件|见附件|同上|待补充|待核对)$",
    re.IGNORECASE,
)
_CARRIER_GENERIC_ONLY_RE = re.compile(
    r"^(?:原件|纸质(?:文件|材料|原件)?|微信(?:聊天记录|聊天)?|聊天记录|电子账单|录屏|导出|文件(?:夹)?|设备|手机|电脑|硬盘|U盘|网盘|账号|原始应用|保存位置|保管位置|路径|目录|档案柜|柜位|卷盒|盒号|案卷号|设备标识|设备编号)$",
    re.IGNORECASE,
)
_CARRIER_GENERIC_PHRASE_RE = re.compile(
    r"^(?:原件文件|原件|纸质文件|纸质材料|微信聊天记录|聊天记录|电子账单|录屏|导出|"
    r"档案柜|柜位|卷盒|盒号|案卷号|设备|设备标识|设备编号|文件夹|文件)"
    r"(?:见材料|见原件|见附件|已核对|已核验|已留存|已保存)?$",
    re.IGNORECASE,
)
_CARRIER_LOCATOR_RE = re.compile(
    # Carrier type words such as ``原件`` or ``纸质`` are not locations by
    # themselves.  Keep only concrete device/file/path/location markers here;
    # a bare carrier type must therefore fall through as missing.
    r"(?:设备|手机|电脑|硬盘|U盘|网盘|账号|原始应用|微信|录屏|文件|文件夹|路径|目录|保存位置|保管位置|档案柜|柜位|卷盒|盒号|案卷号|设备标识|设备编号|电子账单|导出|\.mp4|\.pdf|\.jpg|\.jpeg|\.png|[A-Za-z]:[\\/]|[/\\])",
    re.IGNORECASE,
)

_VIDEO_SOURCE_KINDS = {
    "original_recording",
    "unverified_local_video",
    "public_demo_copy",
    "social_media_copy",
}
_NON_ORIGINAL_VIDEO_SOURCE_KINDS = _VIDEO_SOURCE_KINDS - {"original_recording"}
_VIDEO_EXTENSIONS = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm"}
_SHA256_RE = re.compile(r"^[0-9a-f]{64}$", re.IGNORECASE)


def _carrier_path_exists(value: Any, base_dir: Path | None) -> bool:
    """Require a v2 carrier path to be a real, non-linked file or directory."""

    if not isinstance(value, str) or not value.strip() or base_dir is None:
        return False
    path = Path(os.path.expandvars(value.strip())).expanduser()
    if not path.is_absolute():
        path = base_dir / path
    try:
        # Resolving first would make a junction/symlink look like an ordinary
        # existing carrier. Reject every link-like component before resolving,
        # including a linked final file and a linked parent directory.
        if first_link_component(path) is not None:
            return False
        resolved = path.resolve()
    except OSError:
        return False
    return resolved.is_file() or resolved.is_dir()


_ORDINARY_MATERIAL_REGION_RE = re.compile(
    r"^(?:全图|整图|图片区域|截图区域|聊天区|聊天记录区|主体区|主体区域|内容区|中央金额区|签署栏|群资料页|附件图)$",
    re.IGNORECASE,
)


def _carrier_locator_text(value: Any, *, allow_region: bool = False) -> str:
    """Return a usable carrier locator, never a generic confirmation note."""

    if not isinstance(value, str):
        return ""
    text = value.strip()
    if not text or _CARRIER_PLACEHOLDER_RE.fullmatch(text):
        return ""
    compact = re.sub(r"[\s，,、:：()（）\[\]【】]", "", text)
    if _CARRIER_GENERIC_ONLY_RE.fullmatch(text) or _CARRIER_GENERIC_PHRASE_RE.fullmatch(compact):
        return ""
    if re.search(r"(?:已记录|已核对|已核验|已留存|已保存)$", text) and not re.search(
        r"(?:[:：]\s*[^，。；;]+|[A-Za-z]:[\\/]|[/\\]|\.[A-Za-z]{2,5})",
        text,
    ):
        return ""
    if not _CARRIER_LOCATOR_RE.search(text):
        return ""
    if allow_region and _ORDINARY_MATERIAL_REGION_RE.fullmatch(text):
        return ""
    return text


def _carrier_value_errors(value: Any, *, strict_v2: bool = False, base_dir: Path | None = None) -> list[str]:
    if not strict_v2:
        return []
    if not isinstance(value, dict):
        return ["v2 原始载体必须是对象"]
    locator_keys = ("carrier_location", "path", "location", "region")
    present = [key for key in locator_keys if key in value]
    if not present:
        return ["缺少 carrier_location/path/location/region 定位字段"]
    errors: list[str] = []
    for key in present:
        field = value.get(key)
        if key == "path":
            if not _carrier_path_exists(field, base_dir):
                errors.append("path 必须是相对 manifest 目录或绝对路径下的实际文件/目录")
        elif not _carrier_locator_text(field, allow_region=True):
            errors.append(f"{key} 不能是空值、占位或泛称")
    return errors


def _carrier_value_is_locatable(
    value: Any,
    *,
    strict_v2: bool = False,
    base_dir: Path | None = None,
) -> bool:
    """Check a material's original carrier without trusting dict non-emptiness."""

    if strict_v2:
        return not _carrier_value_errors(value, strict_v2=True, base_dir=base_dir)
    if isinstance(value, dict):
        for key in ("carrier_location", "path", "location", "region"):
            if _carrier_locator_text(value.get(key), allow_region=True):
                return True
        return False
    return bool(_carrier_locator_text(value))


def _material_carrier_value(
    material: dict[str, Any],
    *,
    strict_v2: bool = False,
    base_dir: Path | None = None,
) -> Any:
    """Prefer a locatable original_carrier/carrier value for display and checks."""

    candidates = [material.get("original_carrier"), material.get("carrier")]
    for candidate in candidates:
        if _carrier_value_is_locatable(candidate, strict_v2=strict_v2, base_dir=base_dir):
            return candidate
    for candidate in candidates:
        if candidate is not None:
            return candidate
    return None


def _original_carrier_value(
    material: dict[str, Any],
    data: dict[str, Any],
    *,
    base_dir: Path | None = None,
) -> Any:
    """Actual original carrier; a provenance ``source`` is not a substitute."""

    # A top-level note is not evidence for every image.  It may be retained in
    # the manifest as context, but each material must carry its own carrier or
    # be covered by the structured original_carrier_checked references below.
    return _material_carrier_value(material, strict_v2=bool(data.get("_strict_v2")), base_dir=base_dir)


def _carrier_ref_location(
    ref: Any,
    *,
    strict_v2: bool = False,
    base_dir: Path | None = None,
) -> str:
    if not isinstance(ref, dict):
        return ""
    locator_keys = ("carrier_location", "path", "location", "region")
    if strict_v2:
        if _carrier_value_errors(ref, strict_v2=True, base_dir=base_dir):
            return ""
    for key in ("carrier_location", "path", "location"):
        if key == "path" and strict_v2 and not _carrier_path_exists(ref.get(key), base_dir):
            return ""
        value = _carrier_locator_text(ref.get(key))
        if value:
            return value
    # ``region`` in an ordinary material ref (usually accompanied by page) is
    # an image region such as ``全图`` or ``聊天区``, not an original-carrier
    # locator.  A carrier-specific region remains supported only when it is
    # not shadowed by ordinary material page/region fields and has an explicit
    # device/file/path/location marker.
    if not any(ref.get(key) is not None for key in ("page", "material_page", "attachment_page")):
        value = _carrier_locator_text(ref.get("region"), allow_region=True)
        if value:
            return value
    return ""


def _carrier_coverage(
    refs: Any,
    missing_ids: set[str],
    material_ids: set[str],
    *,
    strict_v2: bool = False,
    base_dir: Path | None = None,
) -> tuple[set[str], list[str]]:
    """Return carrier-covered ids and reject unstructured/unknown refs.

    A carrier ref for a material lacking its own ``original_carrier`` is a
    structured locator, not a free-form confirmation note.  Requiring the
    locator for each missing material prevents a ``note`` or
    ``original_carrier`` placeholder from silently satisfying the formal
    source-carrier gate; ordinary material_refs remain acceptable for
    materials whose carrier is already recorded on the material itself.
    """

    covered: set[str] = set()
    errors: list[str] = []
    for index, ref in enumerate(_normalise_ref_list(refs)):
        if not isinstance(ref, dict):
            errors.append(f"original_carrier_checked 的第{index + 1}个 ref 必须是对象，并写明 material_id 与载体位置")
            continue
        raw_material_id = ref.get("material_id")
        if strict_v2 and (not isinstance(raw_material_id, str) or not raw_material_id.strip()):
            errors.append(f"原始载体 ref {index + 1} 的 material_id 必须是非空字符串")
        material_id = _text(raw_material_id).strip()
        if not material_id:
            errors.append(f"原始载体 ref {index + 1} 缺少 material_id")
            continue
        if material_id not in material_ids:
            errors.append(f"原始载体 ref 引用了未列出的 material_id：{material_id}")
            continue
        if material_id in missing_ids:
            location = _carrier_ref_location(ref, strict_v2=strict_v2, base_dir=base_dir)
            if not location:
                errors.append(f"原始载体 ref {material_id} 缺少可定位的 carrier_location/path/location/region，path 必须实际存在")
            else:
                covered.add(material_id)
    return covered, errors


_MONTH_DAY_TEXT_RE = re.compile(
    r"(?<!\d)(?:0?[1-9]|1[0-2])\s*月\s*(?:0?[1-9]|[12]\d|3[01])\s*[日号]?"
)
_MONTH_DAY_NUMERIC_RE = re.compile(
    r"(?<![\d/.-])(?:0?[1-9]|1[0-2])\s*[/.-]\s*(?:0?[1-9]|[12]\d|3[01])(?![\d/.-])"
)
_RELATIVE_DATE_RE = re.compile(
    r"(?:周[一二三四五六日天]|星期[一二三四五六日天]|礼拜[一二三四五六日天]|今天|昨天|前天|明天|后天|今日|昨日|前日|次日)"
)
_FULL_DATE_RE = re.compile(
    r"(?:19|20)\d{2}\s*年\s*(?:0?[1-9]|1[0-2])\s*月\s*(?:0?[1-9]|[12]\d|3[01])\s*[日号]?"
    r"|(?:19|20)\d{2}\s*[-/.]\s*(?:0?[1-9]|1[0-2])\s*[-/.]\s*(?:0?[1-9]|[12]\d|3[01])"
)


def _date_completeness_messages(text: str, label: str) -> list[str]:
    """Block incomplete calendar expressions while allowing a nearby full date."""

    messages: list[str] = []

    def has_full_date_near(start: int, end: int) -> bool:
        window = text[max(0, start - 24) : min(len(text), end + 24)]
        return bool(_FULL_DATE_RE.search(window))

    for pattern, description in (
        (_MONTH_DAY_TEXT_RE, "缺少年份的月日"),
        (_MONTH_DAY_NUMERIC_RE, "缺少年份的数字月日"),
        (_RELATIVE_DATE_RE, "相对日期或星期表达"),
    ):
        for match in pattern.finditer(text):
            if not has_full_date_near(match.start(), match.end()):
                messages.append(f"{label}含{description}“{match.group(0)}”，正式版需补足明确年份和完整日期")
    return messages


def _normalise_ref_list(value: Any) -> list[Any]:
    if not isinstance(value, list):
        return []
    normalized: list[Any] = []
    for ref in value:
        if isinstance(ref, str):
            normalized.append({"material_id": ref})
        elif isinstance(ref, dict):
            normalized.append(copy.deepcopy(ref))
        else:
            normalized.append(ref)
    return normalized


def _sentence_bounds(text: str, start: int, end: int) -> tuple[int, int]:
    left = max((text.rfind(mark, 0, start) for mark in "。；\n!?！？"), default=-1) + 1
    right_candidates = [text.find(mark, end) for mark in "。；\n!?！？" if text.find(mark, end) >= 0]
    right = min(right_candidates) if right_candidates else len(text)
    return left, right


def _clause_bounds(text: str, sentence_left: int, sentence_right: int, start: int, end: int) -> tuple[int, int]:
    # Split before deciding whether a hit is negated.  A negation in the
    # preceding clause (for example ``不能证明身份，但直接证明真实性``) must
    # not suppress the dangerous conclusion in the following clause.
    separators = list(re.finditer(r"但是|然而|并且|同时|而且|但|且|而|却|仍然|仍|[，,；;:：]", text[sentence_left:sentence_right]))
    clause_left = sentence_left
    clause_right = sentence_right
    for separator in separators:
        absolute_start = sentence_left + separator.start()
        absolute_end = sentence_left + separator.end()
        if absolute_end <= start:
            clause_left = absolute_end
        elif absolute_start >= end:
            clause_right = absolute_start
            break
    return clause_left, clause_right


def _is_negated(text: str, start: int, end: int | None = None) -> bool:
    """Recognise a real boundary clause, not any distant ``不`` character.

    A proof boundary often contains words such as ``承认`` or ``实际支付`` in
    order to say that the image does *not* establish them.  The old twelve-
    character window both missed cross-clause positives and flagged these safe
    boundary sentences.  We therefore inspect the local clause and the
    sentence's forward qualification separately.
    """

    end = start if end is None else end
    sentence_left, sentence_right = _sentence_bounds(text, start, end)
    clause_left, clause_right = _clause_bounds(text, sentence_left, sentence_right, start, end)
    prefix = text[max(clause_left, start - 32) : start]
    direct_negation = re.search(
        r"(?:不能据此(?:认定|证明)?|不得据此(?:认定|证明)?|不据此(?:认定|证明)?|"
        r"不能|不得|无法|未能|并非|不应|不可|不当然|不确认|不等于|不构成|未证明|仅能|仅为|只是|"
        r"属于(?:当事人|申请人)?(?:的)?(?:陈述|主张)?)"
        r"[^，,；;。\n!?！？]{0,24}$",
        prefix,
    )
    if direct_negation:
        return True
    # For ``承认`` and similar words, the qualification frequently follows
    # the hit: ``承认仅为当事人陈述，不能据此认定``.
    matched = text[start:end]
    # Forward qualifications must also remain in the same clause.  A
    # following ``但/且`` clause is a new assertion, not a qualification of
    # the matched risk phrase.
    forward = text[end:clause_right]
    if matched in {"承认", "自认"} and re.search(
        r"(?:仅为|只是|属于(?:当事人|申请人)?的?)(?:事实)?(?:陈述|主张|记录)|不能据此|不得据此|不等于",
        forward,
    ):
        return True
    # A boundary qualifier may be placed after a comma (``承认，仅为当事人
    # 陈述``).  Permit that narrow safe form, but never let a following
    # adversative clause (``承认，但……``) qualify the preceding hit.
    sentence_forward = text[end:sentence_right]
    if matched in {"承认", "自认"} and not re.match(
        r"^[，,；;:\s]*(?:但|然而|却|而|同时|并且)", sentence_forward
    ) and re.match(
        r"^[，,；;:\s]*(?:仅为|只是)(?:(?:当事人|申请人)?的?)(?:事实)?(?:陈述|主张|记录)",
        sentence_forward,
    ):
        return True
    clause = text[clause_left:clause_right]
    if re.search(
        r"(?:不能据此(?:认定|证明)?|不得据此(?:认定|证明)?|不据此(?:认定|证明)?|"
        r"不能|不得|无法|未能|并非|不应|不可|不当然|不确认|不等于|不构成|未证明|仅能|仅为|只是|仅证明|仅载明)",
        clause[: max(0, start - clause_left)],
    ):
        return True
    return False


def _risk_messages(text: str, label: str) -> list[str]:
    messages: list[str] = []
    for name, pattern in RISK_PATTERNS:
        for match in pattern.finditer(text):
            snippet = match.group(0)
            if not _is_negated(text, match.start(), match.end()) and not any(token in snippet for token in ("不能", "不得", "无法", "未能", "并非", "不应", "不可", "未证明", "仅证明", "仅载明", "仅为", "不等于", "不确认", "不证明")):
                messages.append(f"{label}含高风险结论措辞“{name}”，需人工改为证据边界表达")
    return messages


def _risk_fields(value: Any, label: str) -> list[str]:
    """Scan all user-authored claim/evidence text, preserving field context."""

    messages: list[str] = []
    if isinstance(value, dict):
        for key, child in value.items():
            if str(key).startswith("_"):
                continue
            child_label = f"{label}.{key}"
            if isinstance(child, str):
                messages.extend(_risk_messages(child, child_label))
            elif isinstance(child, (dict, list)):
                messages.extend(_risk_fields(child, child_label))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            messages.extend(_risk_fields(child, f"{label}[{index}]"))
    elif isinstance(value, str):
        messages.extend(_risk_messages(value, label))
    return messages


def _numeric_amount(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", "").replace("，", "")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group(0))
    if "万" in text:
        number *= 10000
    elif "千" in text:
        number *= 1000
    return number


def _amount_basis(text: Any) -> str | None:
    value = _text(text).lower().replace(" ", "")
    if any(token in value for token in (
        "订单额", "订单金额", "order_amount", "orderamount", "order_or_agreement_amount",
        "agreement_amount", "contract_amount", "协议额", "协议金额", "合同额", "合同金额",
    )):
        return "order"
    if any(token in value for token in (
        "二维码额", "申请额", "申请金额", "request_amount", "requested_amount", "application_amount",
        "applicationamount", "application_amount_value",
    )):
        return "requested"
    if any(token in value for token in (
        "实际支付", "实际付款", "实付", "实际支付额", "actual_paid", "actual_paid_amount",
        "actual_payment", "actual_payment_amount", "paid_amount", "actualpayment", "actualpaid",
    )):
        return "actual_paid"
    if any(token in value for token in (
        "对方认可", "认可额", "结算额", "结算金额", "acknowledged", "acknowledged_amount",
        "recognized", "recognized_amount", "settled_amount", "settlement_amount", "settlementamount",
    )):
        return "acknowledged"
    return None


_AMOUNT_BASIS_PATTERNS = {
    "order": re.compile(r"(?:订单额|订单金额|order_amount|orderamount|order_or_agreement_amount|agreement_amount|contract_amount|协议额|协议金额|合同额|合同金额)", re.IGNORECASE),
    "requested": re.compile(r"(?:二维码额|申请额|申请金额|request_amount|requested_amount|application_amount|applicationamount|application_amount_value)", re.IGNORECASE),
    "actual_paid": re.compile(r"(?:实际支付|实际付款|实付|实际支付额|actual_paid|actual_paid_amount|actual_payment|actual_payment_amount|paid_amount|actualpayment|actualpaid)", re.IGNORECASE),
    "acknowledged": re.compile(r"(?:对方认可|认可额|结算额|结算金额|acknowledged|acknowledged_amount|recognized|recognized_amount|settled_amount|settlement_amount|settlementamount)", re.IGNORECASE),
}
_AMOUNT_EQUIVALENCE_RE = re.compile(r"(?:即(?:为|是)?|就是|等于|视为|等同于?|相当于|认定为|写成|作为)")


def _amount_bases_in_text(value: Any) -> set[str]:
    text = _text(value).lower().replace(" ", "")
    return {basis for basis, pattern in _AMOUNT_BASIS_PATTERNS.items() if pattern.search(text)}


def _amount_text_equates_bases(value: Any) -> bool:
    text = _text(value)
    for match in _AMOUNT_EQUIVALENCE_RE.finditer(text):
        prefix = text[max(0, match.start() - 4) : match.start()]
        if re.search(r"(?:不|非|未|并非|不能)$", prefix):
            continue
        window = text[max(0, match.start() - 28) : min(len(text), match.end() + 28)]
        if len(_amount_bases_in_text(window)) >= 2:
            return True
    return False


def _amount_payment_blockers(data: dict[str, Any]) -> list[str]:
    """Detect explicit four-basis mismatches and unresolved payment status."""

    blockers: list[str] = []
    explicit_conflicts: list[str] = []
    payment_conflicts: list[str] = []

    def status_tokens(value: Any) -> set[str]:
        """Collect payment-state tokens recursively from a status record.

        ``{"status": "欠付"}`` and ``{"left": {"state": "paid"},
        "right": {"state": "unpaid"}}`` are both status evidence.  The
        previous implementation only inspected children whose *key* already
        contained ``payment`` and therefore missed the nested forms.
        """

        states: set[str] = set()
        if isinstance(value, dict):
            for child in value.values():
                states.update(status_tokens(child))
            return states
        if isinstance(value, list):
            for child in value:
                states.update(status_tokens(child))
            return states
        text = _text(value).lower().replace(" ", "").replace("_", "")
        if not text:
            return states
        unpaid_tokens = (
            "unpaid", "notpaid", "未付款", "未支付", "欠付", "未清偿", "部分付款",
            "partiallypaid", "partialpaid", "部分支付", "尚未清偿",
        )
        unresolved_tokens = (
            "unresolved", "unknown", "disputed", "pending", "待核", "未核",
            "冲突", "不一致", "待确认",
        )
        has_unpaid = any(token in text for token in unpaid_tokens)
        if has_unpaid:
            states.add("unpaid")
        elif any(token in text for token in ("fullypaid", "已清偿", "已付款", "已支付", "已付清", "paid")):
            states.add("paid")
        if any(token in text for token in unresolved_tokens):
            states.add("unresolved")
        return states

    def conflict_is_unresolved(value: Any, field_name: str) -> bool:
        if isinstance(value, list):
            return bool(value)
        if isinstance(value, dict):
            status = value.get("status")
            if status is not None and _text(status).lower() in {"resolved", "clear", "none", "false", "0"}:
                return False
            if value.get("resolved") is True and not value.get("unresolved"):
                return False
            return bool(value) and True
        if isinstance(value, bool):
            return value
        text = _text(value).lower()
        return bool(text) and text not in {"resolved", "clear", "none", "false", "0"}

    payment_key_markers = (
        "payment", "paid", "pay", "付款", "支付", "清偿", "结算", "settlement",
        "欠付", "未清偿", "付款状态", "支付状态", "结算状态",
    )
    status_key_names = {
        "status", "state", "value", "paymentstatus", "paymentstate",
        "settlementstatus", "付款状态", "支付状态", "结算状态",
    }
    record_key_names = {
        "record", "records", "paymentrecord", "paymentrecords", "记录", "双方",
        "left", "right", "debtor", "creditor", "payer", "payee", "side", "sides",
    }
    payment_state_seen: set[tuple[str, tuple[str, ...]]] = set()

    def key_has_payment_context(key: str) -> bool:
        lowered = key.lower().replace("_", "").replace("-", "")
        return any(marker in lowered or marker in key for marker in payment_key_markers)

    def key_is_status_or_record(key: str) -> bool:
        lowered = key.lower().replace("_", "").replace("-", "")
        return lowered in status_key_names or key in status_key_names or lowered in record_key_names or key in record_key_names

    def record_states(label: str, value: Any) -> None:
        states = status_tokens(value)
        if not states:
            return
        state_key = (label, tuple(sorted(states)))
        if state_key in payment_state_seen:
            return
        payment_state_seen.add(state_key)
        # A known unpaid state is a case fact, not a conflict.  Only mixed or
        # unresolved states are blockers, and entity-scoped records are judged
        # independently below.
        if len(states) > 1 or "unresolved" in states:
            payment_conflicts.append(f"{label}={_text(value)}")

    def has_entity_scoped_records(value: Any) -> bool:
        return (
            isinstance(value, list)
            and bool(value)
            and all(
                isinstance(item, dict)
                and _nonempty(item.get("entity_ref"))
                and _nonempty(item.get("time_range") or item.get("payment_time_range"))
                for item in value
            )
        )

    def visit(value: Any, label: str = "manifest", payment_context: bool = False) -> None:
        if isinstance(value, dict):
            # A flat object such as {"order_amount": 100, "actual_paid_amount": 80}
            # is common in v2 fixtures, so classify key/value pairs directly.
            for key, child in value.items():
                key_text = _text(key)
                if label == "manifest" and key_text == "intake_clarification":
                    continue
                if label == "manifest" and key_text.startswith("_"):
                    continue
                lowered = key_text.lower()
                child_text_all = _text(child)
                if lowered.replace("_", "") in {"paymentrecords", "refundrecords"} and isinstance(child, list):
                    for record_index, record in enumerate(child, start=1):
                        record_label = f"{label}.{key_text}[{record_index - 1}]"
                        if not isinstance(record, dict):
                            payment_conflicts.append(f"{record_label} 必须是对象")
                            continue
                        if not _nonempty(record.get("entity_ref")):
                            payment_conflicts.append(f"{record_label} 缺少 entity_ref")
                        if not _nonempty(record.get("time_range") or record.get("payment_time_range")):
                            payment_conflicts.append(f"{record_label} 缺少具体时间范围")
                        states = status_tokens(record.get("status", record.get("state")))
                        if len(states) > 1 or "unresolved" in states:
                            payment_conflicts.append(f"{record_label}={_text(record)}")
                if any(token in child_text_all.lower() for token in ("金额冲突", "金额不一致", "付款冲突", "付款状态冲突", "amount conflict", "payment conflict")):
                    explicit_conflicts.append(f"{label}.{key_text}={child_text_all}")
                if any(token in lowered for token in ("conflict", "冲突", "不一致", "差额")):
                    if isinstance(child, bool) and child or child_text_all.lower() in {"true", "yes", "1", "conflict", "unresolved", "待核", "存在"} or "冲突" in child_text_all or "不一致" in child_text_all:
                        explicit_conflicts.append(f"{label}.{key_text}={child_text_all}")
                if lowered in {"amount_conflicts", "amount_conflict", "unresolved_conflicts", "payment_conflicts", "payment_status_conflicts"} or any(token in lowered for token in ("amount_conflict", "payment_conflict")):
                    if conflict_is_unresolved(child, key_text):
                        explicit_conflicts.append(f"{label}.{key_text}={child_text_all}")
                combined_amount_text = f"{key_text}：{child_text_all}"
                if _amount_text_equates_bases(combined_amount_text):
                    explicit_conflicts.append(f"{label}.{key_text} 含不同金额口径等同表述：{child_text_all}")
                child_payment_context = payment_context or key_has_payment_context(key_text) or key_is_status_or_record(key_text)
                if child_payment_context and not has_entity_scoped_records(child):
                    record_states(f"{label}.{key_text}", child)
                if isinstance(child, dict):
                    descriptor = " ".join(_text(child.get(k, "")) for k in ("basis", "amount_basis", "amount_role", "role", "mode", "label", "name", "type"))
                    child_basis = _amount_basis(descriptor)
                    child_value = next((child.get(k) for k in ("value", "amount", "numeric_value", "金额", "数额") if k in child), None)
                    child_numeric = _numeric_amount(child_value)
                visit(child, f"{label}.{key_text}", child_payment_context)
        elif isinstance(value, list):
            if payment_context and not has_entity_scoped_records(value):
                record_states(label, value)
            for index, child in enumerate(value):
                visit(child, f"{label}[{index}]", payment_context)
        elif payment_context:
            record_states(label, value)
        if isinstance(value, str) and _amount_text_equates_bases(value):
            explicit_conflicts.append(f"{label} 含不同金额口径等同表述：{value}")

    visit(data)
    if explicit_conflicts:
        blockers.append("材料或 Manifest 明确记录金额冲突/不一致：" + "；".join(explicit_conflicts[:5]))
    if payment_conflicts:
        blockers.append("付款状态存在冲突或未核对：" + "；".join(payment_conflicts[:5]))
    return blockers


def _check_refs(
    refs: Any,
    material_ids: set[str],
    label: str,
    blockers: list[str],
    require_location: bool = True,
    *,
    strict_v2: bool = False,
) -> list[dict[str, Any]]:
    normalized_refs = _normalise_ref_list(refs)
    if not normalized_refs:
        blockers.append(f"{label}缺少 material_refs，无法回指材料")
        return []
    output: list[dict[str, Any]] = []
    for ref in normalized_refs:
        if not isinstance(ref, dict):
            blockers.append(f"{label}的 material_refs 必须是对象")
            continue
        raw_material_id = ref.get("material_id")
        if strict_v2 and (not isinstance(raw_material_id, str) or not raw_material_id.strip()):
            blockers.append(f"{label}的 material_id 必须是非空字符串")
        material_id = _text(raw_material_id) if isinstance(raw_material_id, str) else _text(raw_material_id)
        if material_id not in material_ids:
            blockers.append(f"{label}回指不存在的材料：{material_id or '（空）'}")
        if strict_v2:
            if "page" in ref:
                page = ref.get("page")
                valid_page = (
                    (type(page) is int and page > 0)
                    or (isinstance(page, str) and bool(page.strip()))
                )
                if not valid_page:
                    blockers.append(f"{label}的 page 必须是正整数或非空字符串页标签：{material_id or '（空）'}")
            elif require_location:
                blockers.append(f"{label}的材料引用缺少 page：{material_id or '（空）'}")
            if "region" in ref:
                region = ref.get("region")
                if not isinstance(region, str) or not region.strip():
                    blockers.append(f"{label}的 region 必须是非空字符串：{material_id or '（空）'}")
            elif require_location:
                blockers.append(f"{label}的材料引用缺少 region：{material_id or '（空）'}")
        else:
            if require_location and not _nonempty(ref.get("page")):
                blockers.append(f"{label}的材料引用缺少 page：{material_id or '（空）'}")
            if require_location and not _nonempty(ref.get("region")):
                blockers.append(f"{label}的材料引用缺少 region：{material_id or '（空）'}")
        output.append(ref)
    return output


def _validate_proof_targets(
    data: dict[str, Any],
    *,
    strict_v2: bool,
    fatal: list[str],
    blockers: list[str],
) -> dict[tuple[str, str], set[str]]:
    """Validate the pleading/defense target registry used by the v19 matrix."""

    raw_targets = data.get("proof_targets")
    if not isinstance(raw_targets, list):
        if strict_v2 and "proof_targets" in data:
            fatal.append("v2 的 proof_targets 必须是数组")
        if strict_v2:
            blockers.append("v19 正式版必须提供非空 proof_targets[]，列明请求/抗辩及其法律要件")
        return {}
    if strict_v2 and not raw_targets:
        blockers.append("v19 正式版必须提供非空 proof_targets[]，列明请求/抗辩及其法律要件")
        return {}

    registry: dict[tuple[str, str], set[str]] = {}
    for index, target in enumerate(raw_targets, start=1):
        label = f"proof_targets 第{index}项"
        if not isinstance(target, dict):
            (fatal if strict_v2 else blockers).append(f"{label}必须是对象")
            continue
        target_kind = target.get("target_kind")
        target_id = target.get("target_id")
        description = target.get("description")
        legal_elements = target.get("legal_elements")
        if target_kind not in PURPOSE_TARGET_KINDS:
            blockers.append(f"{label}.target_kind 无效：{target_kind!r}")
        if not isinstance(target_id, str) or not target_id.strip():
            blockers.append(f"{label}.target_id 必须是非空字符串")
            continue
        if not isinstance(description, str) or not description.strip():
            blockers.append(f"{label}.description 必须说明具体请求、抗辩、反驳或背景用途")
        if not isinstance(legal_elements, list) or not legal_elements:
            blockers.append(f"{label}.legal_elements 必须是非空字符串数组")
            elements: list[str] = []
        else:
            elements = [item.strip() for item in legal_elements if isinstance(item, str) and item.strip()]
            if len(elements) != len(legal_elements):
                blockers.append(f"{label}.legal_elements 只能包含非空字符串")
            if len(elements) != len(set(elements)):
                blockers.append(f"{label}.legal_elements 不得重复")
        key = (_text(target_kind), target_id.strip())
        if key in registry:
            blockers.append(f"proof_targets 的目标重复：{key[0]}/{key[1]}")
        registry[key] = set(elements)
    return registry


def build_proof_coverage_matrix(data: dict[str, Any]) -> dict[str, Any]:
    """Build a privacy-bounded request-element-evidence coverage sidecar."""

    groups = [item for item in _as_list(data.get("evidence_groups")) if isinstance(item, dict)]
    claims = [item for item in _as_list(data.get("_all_claims")) if isinstance(item, dict)]
    if not claims:
        claims = [item for item in _as_list(data.get("proof_claims")) if isinstance(item, dict)]
        for group in groups:
            claims.extend(item for item in _as_list(group.get("proof_claims")) if isinstance(item, dict))

    all_material_ids: list[str] = []
    group_conflicts: dict[str, list[str]] = {}
    for group in groups:
        group_id = _text(group.get("group_id"))
        all_material_ids.extend(
            _text(material.get("material_id"))
            for material in _as_list(group.get("materials"))
            if isinstance(material, dict) and _text(material.get("material_id"))
        )
        conflicts = [_text(item) for item in _as_list(group.get("conflicts")) if _nonempty(item)]
        if conflicts:
            group_conflicts[group_id] = conflicts

    referenced_with_purpose: set[str] = set()
    claims_without_target: list[str] = []
    cross_group_links: list[dict[str, Any]] = []
    for claim in claims:
        claim_id = _text(claim.get("claim_id"))
        if _nonempty(claim.get("purpose")):
            referenced_with_purpose.update(
                _text(ref.get("material_id"))
                for ref in _as_list(claim.get("material_refs"))
                if isinstance(ref, dict) and _text(ref.get("material_id"))
            )
        basis = claim.get("purpose_basis")
        if not isinstance(basis, dict) or not all(_nonempty(basis.get(key)) for key in ("target_kind", "target_id", "legal_element")):
            if claim_id:
                claims_without_target.append(claim_id)
        if claim.get("_cross_group") or claim.get("cross_group") is True:
            cross_group_links.append(
                {
                    "claim_id": claim_id,
                    "group_ids": list(claim.get("_cross_group_groups", [])),
                    "role": _text(claim.get("role")),
                }
            )

    rows: list[dict[str, Any]] = []
    for target in _as_list(data.get("proof_targets")):
        if not isinstance(target, dict):
            continue
        target_kind = _text(target.get("target_kind"))
        target_id = _text(target.get("target_id"))
        description = _text(target.get("description"))
        for legal_element in target.get("legal_elements", []) if isinstance(target.get("legal_elements"), list) else []:
            if not isinstance(legal_element, str) or not legal_element.strip():
                continue
            matching = []
            for claim in claims:
                basis = claim.get("purpose_basis")
                if not isinstance(basis, dict):
                    continue
                if (
                    _text(basis.get("target_kind")) == target_kind
                    and _text(basis.get("target_id")) == target_id
                    and _text(basis.get("legal_element")) == legal_element.strip()
                ):
                    matching.append(claim)
            row_group_ids = sorted(
                {
                    _text(claim.get("_render_group_id") or claim.get("group_id") or claim.get("_container_group_id"))
                    for claim in matching
                    if _text(claim.get("_render_group_id") or claim.get("group_id") or claim.get("_container_group_id"))
                }
            )
            row_conflicts = [item for gid in row_group_ids for item in group_conflicts.get(gid, [])]
            row_conflicts.extend(
                _text(item)
                for claim in matching
                for item in _as_list(claim.get("conflicts"))
                if _nonempty(item)
            )
            if not matching:
                status = "uncovered"
            elif row_conflicts:
                status = "conflict"
            elif any(
                claim.get("fact_level") in {"element_fact", "procedural_fact"}
                and claim.get("role") in {"direct", "rebuttal"}
                for claim in matching
            ):
                status = "covered"
            else:
                status = "partial"
            rows.append(
                {
                    "target_kind": target_kind,
                    "target_id": target_id,
                    "target_description": description,
                    "legal_element": legal_element.strip(),
                    "status": status,
                    "claim_ids": [_text(claim.get("claim_id")) for claim in matching],
                    "group_ids": row_group_ids,
                    "material_refs": [
                        {
                            "material_id": _text(ref.get("material_id")),
                            "page": ref.get("page"),
                            "region": _text(ref.get("region")),
                        }
                        for claim in matching
                        for ref in _as_list(claim.get("material_refs"))
                        if isinstance(ref, dict)
                    ],
                    "fact_levels": sorted({_text(claim.get("fact_level")) for claim in matching if _text(claim.get("fact_level"))}),
                    "roles": sorted({_text(claim.get("role")) for claim in matching if _text(claim.get("role"))}),
                    "conflicts": list(dict.fromkeys(row_conflicts)),
                }
            )

    unresolved_conflicts = [
        {"scope": "bundle", "text": _text(item)}
        for item in _as_list(data.get("conflicts"))
        if _nonempty(item)
    ]
    unresolved_conflicts.extend(
        {"scope": group_id, "text": text}
        for group_id, items in group_conflicts.items()
        for text in items
    )
    unresolved_conflicts.extend(
        {"scope": "financial", "text": text}
        for text in _amount_payment_blockers(data)
        if _nonempty(text)
    )
    validation = data.get("validation")
    if isinstance(validation, dict):
        unresolved_conflicts.extend(
            {"scope": "validation", "text": _text(item)}
            for item in _as_list(validation.get("unresolved_issues"))
            if _nonempty(item)
        )
    unresolved_conflicts = list(
        {
            (_text(item.get("scope")), _text(item.get("text"))): item
            for item in unresolved_conflicts
            if isinstance(item, dict) and _nonempty(item.get("text"))
        }.values()
    )
    statuses = {row["status"] for row in rows}
    overall_status = (
        "conflict" if unresolved_conflicts or "conflict" in statuses
        else "gap" if "uncovered" in statuses or claims_without_target or set(all_material_ids) - referenced_with_purpose
        else "partial" if "partial" in statuses
        else "covered"
    )
    return {
        "schema_version": 1,
        "model": "request-element-fact-evidence-v19",
        "status": overall_status,
        "coverage": rows,
        "materials_without_purpose": sorted(set(all_material_ids) - referenced_with_purpose),
        "claims_without_target": sorted(set(claims_without_target)),
        "cross_group_links": cross_group_links,
        "unresolved_conflicts": unresolved_conflicts,
    }


def _prepare(data: dict[str, Any], base_dir: Path, manifest_path: Path | None = None) -> tuple[dict[str, Any], list[str], list[str], list[str]]:
    normalized = normalise_manifest(data)
    fatal: list[str] = []
    blockers: list[str] = []
    warnings: list[str] = []
    office_fatal, office_blockers, office_warnings = validate_office_provenance(normalized, base_dir)
    fatal.extend(office_fatal)
    blockers.extend(office_blockers)
    warnings.extend(office_warnings)
    strict_v2 = bool(normalized.get("_strict_v2"))
    if strict_v2:
        if normalized.get("_invalid_evidence_groups_type"):
            fatal.append("v2 的 evidence_groups 必须是数组")
        if normalized.get("_invalid_validation_type"):
            fatal.append("v2 的 validation 必须是对象")
        if normalized.get("_invalid_checks_type"):
            fatal.append("v2 的 validation.checks 必须是对象")
        if normalized.get("_invalid_unresolved_issues_type"):
            fatal.append("v2 的 validation.unresolved_issues 必须是数组，不能静默清空")
        if not isinstance(normalized.get("case_info"), str) or not normalized.get("case_info", "").strip():
            fatal.append("v2 的 case_info 必须是非空字符串")
        if not isinstance(normalized.get("submitter"), str) or not normalized.get("submitter", "").strip():
            fatal.append("v2 的 submitter 必须是非空字符串")
    raw_case_info = normalized.get("case_info")
    if isinstance(raw_case_info, str) and CASE_INFO_PLACEHOLDER.search(raw_case_info):
        blockers.append("case_info 仍含案号或案件信息待补/暂缺/未提供占位；只能生成内部核对稿")
    intake_fatal, intake_blockers = _intake_clarification_gate(normalized)
    fatal.extend(intake_fatal)
    blockers.extend(intake_blockers)
    groups = normalized.get("evidence_groups", [])
    if not groups:
        fatal.append("evidence_groups 必须是非空数组")
    numbers = [group.get("number") for group in groups if isinstance(group, dict)]
    expected_numbers = list(range(1, len(groups) + 1))
    if numbers != expected_numbers:
        fatal.append(f"编号必须从1连续递增，当前为：{numbers}")

    # Context materials are retained as provenance/background only.  They are
    # intentionally kept outside the evidence material set so they cannot be
    # embedded, counted in the filing table, or used as claim evidence.
    context_materials = normalized.get("context_materials", [])
    context_ids: set[str] = set()
    if normalized.get("_invalid_context_materials_type"):
        fatal.append("v2 的 context_materials 必须是数组")
    if not isinstance(context_materials, list):
        context_materials = []
    for index, context in enumerate(context_materials, start=1):
        if not isinstance(context, dict):
            (fatal if strict_v2 else blockers).append(f"context_materials 第{index}项必须是对象")
            continue
        raw_context_id = context.get("material_id")
        if strict_v2 and (not isinstance(raw_context_id, str) or not raw_context_id.strip()):
            fatal.append(f"context_materials 第{index}项的 material_id 必须是非空字符串")
        context_id = _text(raw_context_id).strip()
        if not context_id:
            (fatal if strict_v2 else blockers).append(f"context_materials 第{index}项缺少 material_id")
            continue
        if context_id in context_ids:
            fatal.append(f"context_materials 的 material_id 重复：{context_id}")
        context_ids.add(context_id)

    media_source_records: dict[str, dict[str, Any]] = {}
    raw_media_sources = normalized.get("media_sources", [])
    if "media_sources" in normalized and not isinstance(raw_media_sources, list):
        fatal.append("v2 的 media_sources 必须是数组")
        raw_media_sources = []
    for source_number, raw_source in enumerate(raw_media_sources, start=1):
        if not isinstance(raw_source, dict):
            fatal.append(f"media_sources 第{source_number}项必须是对象")
            continue
        raw_source_id = raw_source.get("source_id")
        source_id = raw_source_id.strip() if isinstance(raw_source_id, str) else ""
        if not source_id:
            fatal.append(f"media_sources 第{source_number}项缺少非空 source_id")
            continue
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}", source_id) or ".." in source_id:
            fatal.append(f"media_sources 第{source_number}项的 source_id 含不安全路径字符")
            continue
        if source_id in media_source_records:
            fatal.append(f"media_sources 的 source_id 重复：{source_id}")
            continue
        source_kind = _text(raw_source.get("source_kind")).strip().lower()
        if source_kind not in _VIDEO_SOURCE_KINDS:
            fatal.append(f"media_sources {source_id} 的 source_kind 无效：{source_kind or '（空）'}")
        try:
            source_path = resolve_material_path(raw_source.get("path"), base_dir)
        except ValueError as exc:
            fatal.append(f"media_sources {source_id} 的视频路径无效：{exc}")
            continue
        if not source_path.is_file():
            fatal.append(f"media_sources {source_id} 的本地视频不存在：{source_path}")
            continue
        declared_digest = raw_source.get("sha256")
        if not isinstance(declared_digest, str) or not _SHA256_RE.fullmatch(declared_digest.strip()):
            fatal.append(f"media_sources {source_id} 的 sha256 必须是64位十六进制字符串")
            source_digest = _sha256(source_path)
        else:
            source_digest = _sha256(source_path)
            if declared_digest.strip().lower() != source_digest:
                fatal.append(f"media_sources {source_id} 的 sha256 与本地视频不一致")
        duration = raw_source.get("duration_seconds")
        if isinstance(duration, bool) or not isinstance(duration, (int, float)) or not math.isfinite(float(duration)) or float(duration) <= 0:
            fatal.append(f"media_sources {source_id} 的 duration_seconds 必须是正数")
        stream_index = raw_source.get("stream_index")
        if isinstance(stream_index, bool) or not isinstance(stream_index, int) or stream_index < 0:
            fatal.append(f"media_sources {source_id} 的 stream_index 必须是非负整数")
        try:
            frame_index_path = resolve_material_path(raw_source.get("frame_index_path"), base_dir)
        except ValueError as exc:
            fatal.append(f"media_sources {source_id} 的 frame_index_path 无效：{exc}")
            frame_index_path = None
        index_frames: dict[Path, dict[str, Any]] = {}
        if frame_index_path is not None:
            if not frame_index_path.is_file():
                fatal.append(f"media_sources {source_id} 的截帧索引不存在：{frame_index_path}")
            else:
                try:
                    frame_index = json.loads(frame_index_path.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError) as exc:
                    fatal.append(f"media_sources {source_id} 的截帧索引无法读取：{exc}")
                    frame_index = {}
                indexed_source = frame_index.get("source_video") if isinstance(frame_index, dict) else None
                if not isinstance(indexed_source, dict):
                    fatal.append(f"media_sources {source_id} 的截帧索引缺少 source_video")
                else:
                    if _text(indexed_source.get("source_id")).strip() != source_id:
                        fatal.append(f"media_sources {source_id} 与截帧索引的 source_id 不一致")
                    if _text(indexed_source.get("source_kind")).strip().lower() != source_kind:
                        fatal.append(f"media_sources {source_id} 与截帧索引的 source_kind 不一致")
                    if indexed_source.get("path"):
                        try:
                            indexed_source_path = resolve_material_path(indexed_source.get("path"), frame_index_path.parent)
                        except ValueError:
                            indexed_source_path = None
                        if indexed_source_path != source_path:
                            fatal.append(f"media_sources {source_id} 与旧版截帧索引回指的原视频不一致")
                    elif (
                        _text(indexed_source.get("filename")) != source_path.name
                        or indexed_source.get("size") != source_path.stat().st_size
                    ):
                        fatal.append(f"media_sources {source_id} 与截帧索引的文件名或大小不一致")
                    if _text(indexed_source.get("sha256")).strip().lower() != source_digest:
                        fatal.append(f"media_sources {source_id} 与截帧索引的原视频摘要不一致")
                    if indexed_source.get("stream_index") != stream_index:
                        fatal.append(f"media_sources {source_id} 与截帧索引的视频流编号不一致")
                indexed_frames = frame_index.get("frames") if isinstance(frame_index, dict) else None
                if not isinstance(indexed_frames, list) or not indexed_frames:
                    fatal.append(f"media_sources {source_id} 的截帧索引缺少 frames")
                else:
                    for indexed_frame in indexed_frames:
                        if not isinstance(indexed_frame, dict):
                            fatal.append(f"media_sources {source_id} 的截帧索引含非对象帧")
                            continue
                        relative_frame = Path(_text(indexed_frame.get("path")))
                        if relative_frame.is_absolute() or ".." in relative_frame.parts or not relative_frame.parts:
                            fatal.append(f"media_sources {source_id} 的截帧索引含不安全帧路径")
                            continue
                        indexed_frame_path = (frame_index_path.parent / relative_frame).resolve()
                        if not indexed_frame_path.is_file():
                            fatal.append(f"media_sources {source_id} 的索引帧不存在：{indexed_frame_path}")
                            continue
                        indexed_digest = _text(indexed_frame.get("sha256")).strip().lower()
                        if not _SHA256_RE.fullmatch(indexed_digest) or indexed_digest != _sha256(indexed_frame_path):
                            fatal.append(f"media_sources {source_id} 的索引帧摘要不一致：{indexed_frame_path.name}")
                            continue
                        index_frames[indexed_frame_path] = indexed_frame
        media_source_records[source_id] = {
            "source_id": source_id,
            "source_kind": source_kind,
            "path": source_path,
            "sha256": source_digest,
            "duration_seconds": duration,
            "stream_index": stream_index,
            "frame_index_path": frame_index_path,
            "frames": index_frames,
        }

    group_ids: set[str] = set()
    material_ids: set[str] = set()
    material_records: dict[str, dict[str, Any]] = {}
    path_hashes: dict[str, list[str]] = {}
    material_paths: list[Path] = []
    missing_carrier_ids: list[str] = []
    last_video_times: dict[str, float] = {}
    manifest_resolved = manifest_path.resolve() if manifest_path else None
    for group in groups:
        if not isinstance(group, dict):
            fatal.append("证据组必须是对象")
            continue
        group_id = _text(group.get("group_id"))
        if strict_v2:
            if isinstance(group.get("number"), bool) or not isinstance(group.get("number"), int):
                fatal.append(f"第{group.get('number', '?')}组的 number 必须是整数")
            for required_key, label in (("group_id", "group_id"), ("evidence_name", "证据名称"), ("evidence_form", "证据形式"), ("proof_object", "proof_object")):
                if not isinstance(group.get(required_key), str) or not group.get(required_key, "").strip():
                    fatal.append(f"第{group.get('number', group_id or '?')}组缺少非空{label}")
            if group.get("_invalid_materials_type"):
                fatal.append(f"第{group.get('number', '?')}组 materials 必须是数组")
            if group.get("_invalid_proof_claims_type"):
                fatal.append(f"第{group.get('number', '?')}组 proof_claims 必须是数组")
        raw_proof_object = group.get("proof_object")
        if isinstance(raw_proof_object, str) and raw_proof_object.strip():
            display_proof_object = _display_proof_object(raw_proof_object)
            if not display_proof_object:
                fatal.append(
                    f"第{group.get('number', group_id or '?')}组 proof_object 仅含核验、补强或证明边界提示；"
                    "请改写为积极、可回指的待证事实，并说明其服务的请求、抗辩或证明要件"
                )
            else:
                group["_display_proof_object"] = display_proof_object
        if group_id in group_ids:
            fatal.append(f"group_id 重复：{group_id}")
        group_ids.add(group_id)
        materials = group.get("materials")
        if not isinstance(materials, list) or not materials:
            fatal.append(f"第{group.get('number', '?')}组 materials 必须为非空数组")
            continue
        for material in materials:
            if not isinstance(material, dict):
                fatal.append(f"第{group.get('number', '?')}组存在非对象材料")
                continue
            raw_material_id = material.get("material_id")
            if strict_v2 and (not isinstance(raw_material_id, str) or not raw_material_id.strip()):
                fatal.append(f"第{group.get('number', '?')}组材料的 material_id 必须是非空字符串")
            material_id = _text(raw_material_id)
            if not material_id:
                fatal.append(f"第{group.get('number', '?')}组材料缺少 material_id")
                continue
            if material_id in material_ids:
                fatal.append(f"material_id 全局重复：{material_id}")
            if material_id in context_ids:
                fatal.append(f"证据材料与 context_materials 共用 material_id：{material_id}")
            material_ids.add(material_id)
            raw_path = material.get("path")
            try:
                path = resolve_material_path(raw_path, base_dir)
            except ValueError as exc:
                fatal.append(f"材料 {material_id} 的路径无效：{exc}")
                continue
            if manifest_resolved and path == manifest_resolved:
                fatal.append(f"材料 {material_id} 的路径不得等于 manifest：{path}")
            if not path.is_file():
                fatal.append(f"材料不存在（{material_id}）：{path}")
                continue
            if path.suffix.lower() not in ALLOWED_IMAGES:
                fatal.append(f"材料 {material_id} 类型不受支持（unsupported image type / 错误类型）：{path}；标准卷仅接受图片（images only）")
                continue
            try:
                display_rotation = _display_rotation(material.get("display_rotation_degrees_clockwise"))
            except ValueError as exc:
                fatal.append(f"材料 {material_id} 的显示旋转字段无效：{exc}")
                display_rotation = 0
            record = _image_preflight(
                path,
                material.get("readability"),
                display_rotation_degrees_clockwise=display_rotation,
            )
            record.update(
                {
                    "material_id": material_id,
                    "group_id": group_id,
                    "source": _source_value(material),
                    "original_carrier": _original_carrier_value(material, normalized, base_dir=base_dir),
                    "source_time": material.get("source_time", ""),
                    "fact": material.get("fact", ""),
                    "path_text": _text(raw_path),
                }
            )
            if record["decode_status"] != "ok":
                fatal.extend(f"材料 {material_id}：{item}" for item in record["errors"])
            else:
                warnings.extend(f"材料 {material_id}：{item}" for item in record["warnings"])
                if record["errors"]:
                    blockers.extend(f"材料 {material_id}：{item}" for item in record["errors"])
            if strict_v2:
                for carrier_key in ("original_carrier", "carrier"):
                    if carrier_key in material:
                        for carrier_error in _carrier_value_errors(
                            material.get(carrier_key),
                            strict_v2=True,
                            base_dir=base_dir,
                        ):
                            blockers.append(f"材料 {material_id} 的 {carrier_key}：{carrier_error}")
            material_carrier = _material_carrier_value(material, strict_v2=strict_v2, base_dir=base_dir)
            if not _carrier_value_is_locatable(material_carrier, strict_v2=strict_v2, base_dir=base_dir):
                missing_carrier_ids.append(material_id)
                if strict_v2:
                    blockers.append(
                        f"材料 {material_id} 的 original_carrier/carrier 缺少可定位的 carrier_location/path/location/region；"
                        "note、description、status 或普通核验 note 不能替代原始载体位置"
                    )
            if strict_v2 and isinstance(material_carrier, dict):
                carrier_source_kind = _text(material_carrier.get("source_kind")).strip().lower()
                carrier_path: Path | None = None
                if _nonempty(material_carrier.get("path")):
                    try:
                        carrier_path = resolve_material_path(material_carrier.get("path"), base_dir)
                    except ValueError:
                        carrier_path = None
                is_video_carrier = (
                    carrier_source_kind in _VIDEO_SOURCE_KINDS
                    or (carrier_path is not None and carrier_path.suffix.lower() in _VIDEO_EXTENSIONS)
                    or "video" in _text(material_carrier.get("carrier_type")).lower()
                    or "录屏" in _text(material_carrier.get("carrier_type"))
                )
                video_frame = material.get("video_frame")
                if is_video_carrier and not isinstance(video_frame, dict):
                    fatal.append(f"材料 {material_id} 的视频截帧缺少结构化 video_frame")
                if isinstance(video_frame, dict):
                    source_id = _text(video_frame.get("source_id")).strip()
                    source_record = media_source_records.get(source_id)
                    if not source_id or source_record is None:
                        fatal.append(f"材料 {material_id} 的 video_frame.source_id 未回指 media_sources")
                    timestamp = video_frame.get("source_time_seconds")
                    if isinstance(timestamp, bool) or not isinstance(timestamp, (int, float)) or not math.isfinite(float(timestamp)) or float(timestamp) < 0:
                        fatal.append(f"材料 {material_id} 的 video_frame.source_time_seconds 必须是非负数")
                        timestamp_value = None
                    else:
                        timestamp_value = float(timestamp)
                    timecode = video_frame.get("source_timecode")
                    if not isinstance(timecode, str) or not timecode.strip():
                        fatal.append(f"材料 {material_id} 的 video_frame.source_timecode 必须是非空字符串")
                    elif _text(material.get("source_time")).strip() != timecode.strip():
                        fatal.append(f"材料 {material_id} 的 source_time 与 video_frame.source_timecode 不一致")
                    source_frame_index = video_frame.get("source_frame_index")
                    if source_frame_index is not None and (
                        isinstance(source_frame_index, bool)
                        or not isinstance(source_frame_index, int)
                        or source_frame_index < 0
                    ):
                        fatal.append(f"材料 {material_id} 的 video_frame.source_frame_index 必须是非负整数或 null")
                    if source_record is not None:
                        source_kind = source_record["source_kind"]
                        if source_kind in _NON_ORIGINAL_VIDEO_SOURCE_KINDS:
                            blockers.append(
                                f"材料 {material_id} 的视频来源类型为 {source_kind}；"
                                "未核实本地视频或公开/社交媒体派生副本不能作为案件原始载体进入正式版"
                            )
                        if carrier_source_kind != source_kind:
                            fatal.append(f"材料 {material_id} 的 original_carrier.source_kind 与 media_sources 不一致")
                        carrier_status = _text(material_carrier.get("status")).strip().lower()
                        if source_kind == "original_recording" and carrier_status in {"unverified", "derived_copy_not_original"}:
                            fatal.append(f"材料 {material_id} 的 original_carrier.status 与原录屏来源类型冲突")
                        if source_kind in _NON_ORIGINAL_VIDEO_SOURCE_KINDS and carrier_status == "candidate_original":
                            fatal.append(f"材料 {material_id} 的 original_carrier.status 错把派生/未核实视频标为原始候选")
                        if carrier_path != source_record["path"]:
                            fatal.append(f"材料 {material_id} 的 original_carrier.path 与 media_sources 原视频不一致")
                        indexed_frame = source_record["frames"].get(path)
                        if indexed_frame is None:
                            fatal.append(f"材料 {material_id} 的图片未出现在对应 video-frame-index.json")
                        else:
                            indexed_timestamp = indexed_frame.get("timestamp_seconds")
                            if timestamp_value is not None and (
                                isinstance(indexed_timestamp, bool)
                                or not isinstance(indexed_timestamp, (int, float))
                                or abs(float(indexed_timestamp) - timestamp_value) > 0.001
                            ):
                                fatal.append(f"材料 {material_id} 的数值时间与截帧索引不一致")
                            if isinstance(timecode, str) and indexed_frame.get("source_timecode") != timecode:
                                fatal.append(f"材料 {material_id} 的格式化时间码与截帧索引不一致")
                            indexed_video_frame = indexed_frame.get("video_frame")
                            if not isinstance(indexed_video_frame, dict) or _text(indexed_video_frame.get("source_id")).strip() != source_id:
                                fatal.append(f"材料 {material_id} 的截帧索引未回指同一 source_id")
                            indexed_frame_index = indexed_frame.get("source_frame_index")
                            if source_frame_index is not None and indexed_frame_index is not None and source_frame_index != indexed_frame_index:
                                fatal.append(f"材料 {material_id} 的 source_frame_index 与截帧索引不一致")
                        if timestamp_value is not None:
                            previous = last_video_times.get(source_id)
                            if previous is not None and timestamp_value < previous:
                                fatal.append(f"材料 {material_id} 未按 source_time_seconds 数值升序排列")
                            last_video_times[source_id] = timestamp_value
                elif carrier_source_kind in _NON_ORIGINAL_VIDEO_SOURCE_KINDS:
                    blockers.append(
                        f"材料 {material_id} 的视频来源类型为 {carrier_source_kind}；"
                        "未核实本地视频或公开/社交媒体派生副本不能作为案件原始载体进入正式版"
                    )
            material_conflicts = material.get("conflicts", [])
            if isinstance(material_conflicts, list):
                blockers.extend(f"材料 {material_id} 存在未解决冲突：{_text(item)}" for item in material_conflicts if _nonempty(item))
            material_records[material_id] = record
            material_paths.append(path)
            if record.get("decode_status") == "ok":
                digest = _sha256(path)
                record["sha256"] = digest
                path_hashes.setdefault(digest, []).append(material_id)

    duplicate_groups: list[dict[str, Any]] = []
    alias_mapping: dict[str, str] = {}
    # Keep canonical declarations and canonical->alias declarations separate.
    # Treating both directions as "canonical candidates" makes a valid
    # bidirectional pair look contradictory (M001->M002 and M002->M001).
    declared_canonical_targets: dict[str, set[str]] = {}
    declared_alias_targets: dict[str, set[str]] = {}
    for _group, material in _all_material_items(normalized):
        if not isinstance(material, dict):
            continue
        material_id = _text(material.get("material_id"))
        declared = material.get("canonical_material_id") or material.get("canonical_id") or material.get("alias_of")
        if _nonempty(declared):
            declared_canonical_targets.setdefault(material_id, set()).add(_text(declared))
        aliases = material.get("material_aliases")
        if isinstance(aliases, list):
            for alias in aliases:
                if _nonempty(alias):
                    declared_alias_targets.setdefault(material_id, set()).add(_text(alias))

    # Validate every declared edge, including aliases named only from a
    # canonical material.  Unknown ids, one-way declarations and different
    # bytes must never be silently ignored before duplicate selection.
    declared_edges: dict[str, set[str]] = {}
    for declared_id, targets in declared_canonical_targets.items():
        declared_edges.setdefault(declared_id, set()).update(targets)
    for canonical_id, aliases in declared_alias_targets.items():
        declared_edges.setdefault(canonical_id, set()).update(aliases)

    for declared_id, targets in declared_edges.items():
        if declared_id not in material_ids:
            blockers.append(f"alias/canonical 声明引用未列出的 material_id：{declared_id}")
            continue
        for target_id in targets:
            if target_id not in material_ids:
                blockers.append(f"材料 {declared_id} 声明了未列出的 alias/canonical：{target_id}")
                continue
            if target_id != declared_id and declared_id not in declared_edges.get(target_id, set()):
                blockers.append(f"alias/canonical 声明不是双向：{declared_id} -> {target_id} 缺少反向声明")
            source_digest = material_records.get(declared_id, {}).get("sha256")
            target_digest = material_records.get(target_id, {}).get("sha256")
            if source_digest and target_digest and source_digest != target_digest:
                blockers.append(f"alias/canonical 声明的 SHA-256 不一致：{declared_id}={source_digest}，{target_id}={target_digest}")
    for digest, ids in path_hashes.items():
        if len(ids) > 1:
            candidates: set[str] = set()
            for material_id in ids:
                # A material explicitly naming a canonical target contributes
                # that target; a canonical material listing aliases contributes
                # itself.  Reverse alias declarations do not make the alias a
                # second canonical candidate.
                candidates.update(
                    target for target in declared_canonical_targets.get(material_id, set())
                    if target in ids
                )
                if declared_alias_targets.get(material_id):
                    candidates.add(material_id)
            if len(candidates) > 1:
                blockers.append(f"SHA-256 重复组 {ids} 的 canonical 声明相互冲突：{sorted(candidates)}")
            canonical = sorted(candidates)[0] if candidates else ids[0]
            if canonical not in ids:
                blockers.append(f"canonical_material_id {canonical} 不在同一 SHA-256 重复组：{ids}")
                canonical = ids[0]
            aliases = ids[1:]
            aliases = [material_id for material_id in ids if material_id != canonical]
            duplicate_groups.append({"sha256": digest, "canonical": canonical, "aliases": aliases})
            for alias in aliases:
                alias_mapping[alias] = canonical
                material_records[alias]["canonical_material_id"] = canonical
            material_records[canonical]["canonical_material_id"] = canonical
        elif ids:
            only_id = ids[0]
            declared = declared_canonical_targets.get(only_id, set())
            if declared and declared != {only_id}:
                blockers.append(f"材料 {only_id} 声明了不存在同 SHA-256 的 canonical：{sorted(declared)}")
            material_records[only_id]["canonical_material_id"] = only_id

    # Conflicts are a formal gate, not a reason to silently choose one side.
    for group in groups:
        conflicts = group.get("conflicts", []) if isinstance(group, dict) else []
        if isinstance(conflicts, list):
            blockers.extend(f"第{group.get('number', '?')}组存在未解决冲突：{_text(item)}" for item in conflicts if _nonempty(item))
    top_conflicts = normalized.get("conflicts", [])
    if isinstance(top_conflicts, list):
        blockers.extend(f"存在未解决冲突：{_text(item)}" for item in top_conflicts if _nonempty(item))
    blockers.extend(_amount_payment_blockers(normalized))

    validation = normalized.get("validation", {})
    if normalized.get("_source_schema_version") == 1:
        warnings.append("检测到 Manifest v1：仅生成待核对草稿，旧版布尔核验不能作为正式版依据")
    if normalized.get("_legacy_boolean_checks"):
        blockers.append("核验字段仍含旧版布尔值；必须改为含 status、material_refs、note 的记录对象")
    if validation.get("status") != "resolved":
        blockers.append("validation.status 未设为 resolved")
    unresolved = validation.get("unresolved_issues", [])
    if not isinstance(unresolved, list):
        fatal.append("validation.unresolved_issues 必须是数组，非法类型不得按空数组处理")
    else:
        blockers.extend(f"存在未解决事项：{_text(item)}" for item in unresolved if _nonempty(item))
    checks = validation.get("checks", {})
    for key in CHECK_KEYS:
        check = checks.get(key, {})
        if not isinstance(check, dict):
            blockers.append(f"核验项 {key} 不是记录对象")
            continue
        status = check.get("status")
        refs = check.get("material_refs")
        note = _text(check.get("note"))
        if status not in CHECK_STATUSES:
            blockers.append(f"核验项 {key} 的 status 无效：{status!r}")
        if status == "unresolved":
            blockers.append(f"核验项 {key} 尚未解决")
        if status == "resolved" and (not _as_list(refs) or not note.strip()):
            blockers.append(f"核验项 {key} 已 resolved 但缺少具体 material_refs 或 note")
        if status == "not_applicable" and not note.strip():
            blockers.append(f"核验项 {key} 为 not_applicable 但缺少 note")
        if _as_list(refs):
            _check_refs(refs, material_ids, f"核验项 {key}", blockers, require_location=False, strict_v2=strict_v2)
    carrier_check = checks.get("original_carrier_checked", {})
    missing_carrier_set = set(missing_carrier_ids)
    carrier_refs = carrier_check.get("material_refs", []) if isinstance(carrier_check, dict) else []
    covered_carriers, carrier_ref_errors = _carrier_coverage(
        carrier_refs,
        missing_carrier_set,
        material_ids,
        strict_v2=strict_v2,
        base_dir=base_dir,
    )
    blockers.extend(carrier_ref_errors)
    carrier_check_evidence = (
        isinstance(carrier_check, dict)
        and carrier_check.get("status") == "resolved"
        and not missing_carrier_set.difference(covered_carriers)
        and not carrier_ref_errors
    )
    for material_id in missing_carrier_ids:
        if material_id not in covered_carriers:
            blockers.append(f"材料 {material_id} 缺少逐项原始载体 original_carrier 或 carrier ref/位置；source 只能说明来源，不能替代载体核验")
    if missing_carrier_set and not carrier_check_evidence:
        blockers.append("original_carrier_checked 未逐材料提供可定位的原始载体依据；不得用全局 note 或其他材料的载体放行")
    if not missing_carrier_set and (not isinstance(carrier_check, dict) or carrier_check.get("status") not in {"resolved", "not_applicable"}):
        blockers.append("所有材料虽有逐项 original_carrier，但 original_carrier_checked 未明确 resolved")

    if strict_v2 and normalized.get("_invalid_top_proof_claims_type"):
        fatal.append("v2 的顶层 proof_claims 必须是数组")
    target_registry = _validate_proof_targets(
        normalized,
        strict_v2=strict_v2,
        fatal=fatal,
        blockers=blockers,
    )
    claims = normalized.get("_all_claims", [])
    claim_ids: set[str] = set()
    material_group_by_id: dict[str, str] = {}
    group_order: dict[str, int] = {}
    group_number_aliases: dict[str, str] = {}
    group_number_by_id: dict[str, str] = {}
    for order, group in enumerate(groups):
        if not isinstance(group, dict):
            continue
        gid = _text(group.get("group_id"))
        group_order[gid] = order
        if _nonempty(group.get("number")):
            group_number_aliases[_text(group.get("number"))] = gid
            group_number_by_id[gid] = _text(group.get("number"))
        for material in group.get("materials", []) if isinstance(group.get("materials"), list) else []:
            if isinstance(material, dict):
                material_group_by_id[_text(material.get("material_id"))] = gid

    for claim in claims:
        if not isinstance(claim, dict):
            if strict_v2:
                fatal.append("v2 的 proof_claims 每一项必须是对象")
            else:
                blockers.append("proof_claims 存在非对象项")
            continue
        raw_claim_id = claim.get("claim_id")
        if strict_v2 and (not isinstance(raw_claim_id, str) or not raw_claim_id.strip()):
            fatal.append("v2 的 proof_claim 的 claim_id 必须是非空字符串；不得用列表、对象或数字替代")
        claim_id = _text(raw_claim_id) if isinstance(raw_claim_id, str) else ""
        label = f"proof_claim {claim_id or '（空）'}"
        if not claim_id:
            blockers.append(f"{label} 缺少必填 claim_id；不得自动生成")
        elif claim_id in claim_ids:
            blockers.append(f"proof_claim claim_id 重复：{claim_id}")
        claim_ids.add(claim_id)
        claim_type = claim.get("type")
        role = claim.get("role")
        fact_level = claim.get("fact_level")
        if claim_type not in {"fact", "legal_significance"}:
            blockers.append(f"{label} 的 type 必须为 fact 或 legal_significance")
        if role not in CLAIM_ROLES:
            blockers.append(
                f"{label} 的 role 不合法；只能是 direct、indirect、corroborative、rebuttal 或 linking"
            )
        if strict_v2 and fact_level not in FACT_LEVELS:
            blockers.append(
                f"{label} 的 fact_level 不合法；只能是 element_fact、indirect_fact、auxiliary_fact 或 procedural_fact"
            )
        if not isinstance(claim.get("text"), str) or not claim.get("text", "").strip():
            (fatal if strict_v2 else blockers).append(f"{label} 的 text 必须是非空字符串")
        if not isinstance(claim.get("boundary"), str) or not claim.get("boundary", "").strip():
            (fatal if strict_v2 else blockers).append(f"{label} 的 boundary 必须是非空字符串")
        purpose = claim.get("purpose")
        if not isinstance(purpose, str) or not purpose.strip():
            blockers.append(f"{label} 缺少 purpose；必须说明提交该材料拟支持或反驳的请求、抗辩或证明要件")
        else:
            purpose_text = purpose.strip()
            purpose_compact = re.sub(r"[\s。；，：:、,.，]+", "", purpose_text)
            content_compacts = {
                re.sub(r"[\s。；，：:、,.，]+", "", _text(claim.get(key)).strip())
                for key in ("text", "atomic_fact", "fact", "legal_significance")
                if _nonempty(claim.get(key))
            }
            if purpose_text in {"证明该材料", "证明内容", "证明事实", "证明作用", "证明相关事实"}:
                blockers.append(f"{label} 的 purpose 过于空泛；必须说明服务的请求、抗辩或证明要件")
            elif purpose_compact and purpose_compact in content_compacts:
                blockers.append(f"{label} 的 purpose 仅复述材料内容；必须说明服务的请求、抗辩或证明要件")
            else:
                purpose_core = _purpose_semantic_core(purpose_text)
                content_cores = {
                    _purpose_semantic_core(claim.get(key))
                    for key in ("text", "atomic_fact", "fact", "legal_significance")
                    if _nonempty(claim.get(key))
                }
                if (
                    purpose_core
                    and len(purpose_core) >= 8
                    and any(
                        core
                        and len(core) >= 8
                        and (core in purpose_core or purpose_core in core)
                        for core in content_cores
                    )
                    and not _purpose_has_concrete_anchor(purpose_text)
                ):
                    blockers.append(f"{label} 的 purpose 与材料内容实质同义复述；必须补充具体请求、抗辩或证明要件用途")
                elif not re.search(r"请求|抗辩|要件|争议|主张|反驳|诉讼|仲裁|权利|义务|返还|赔偿|履行", purpose_text):
                    blockers.append(f"{label} 的 purpose 未对应具体请求、抗辩、争议焦点或证明要件")
                elif not (
                    _target_clause_has_specific_content(purpose_text)
                    or _purpose_has_concrete_anchor(purpose_text)
                ):
                    blockers.append(
                        f"{label} 的 purpose 过于空泛；不能只写本案请求、法律要件等通用标签，"
                        "必须写明具体请求内容、抗辩理由或法律要件"
                    )
        purpose_basis = claim.get("purpose_basis")
        if not isinstance(purpose_basis, dict):
            blockers.append(f"{label} 缺少 purpose_basis 对象；必须记录用途依据及核验状态")
        else:
            purpose_kind = purpose_basis.get("kind")
            purpose_reference = purpose_basis.get("reference")
            purpose_status = purpose_basis.get("status")
            if purpose_kind not in PURPOSE_BASIS_KINDS:
                blockers.append(
                    f"{label}.purpose_basis.kind 无效：{purpose_kind!r}；"
                    "只能是 pleading、user_statement 或 materials_only"
                )
            if not isinstance(purpose_reference, str) or not purpose_reference.strip():
                blockers.append(f"{label}.purpose_basis.reference 必须是非空字符串")
            elif purpose_kind == "pleading" and not _pleading_reference_has_locator(purpose_reference):
                blockers.append(
                    f"{label}.purpose_basis.reference 必须定位到书状具体页码、段落、条项或明确章节；"
                    "仅写起诉状、答辩状或仲裁申请书名称不足以复核"
                )
            if purpose_status not in PURPOSE_BASIS_STATUSES:
                blockers.append(
                    f"{label}.purpose_basis.status 无效：{purpose_status!r}；"
                    "只能是 confirmed 或 provisional"
                )
            elif purpose_status == "provisional":
                blockers.append(f"{label}.purpose_basis 仍为 provisional；用途尚待用户、律师或相关书状确认，只能生成草稿")
            if purpose_kind == "materials_only":
                blockers.append(f"{label}.purpose_basis 仅依据材料归纳（materials_only）；不得据此生成正式版")
                if purpose_status != "provisional":
                    blockers.append(
                        f"{label}.purpose_basis 的 materials_only 不得标记为 confirmed；"
                        "必须保持 provisional 并在草稿中披露缺失"
                    )
            if purpose_kind == "user_statement" and purpose_status == "confirmed" and isinstance(purpose_reference, str):
                compact_reference = re.sub(r"[\s。；，：:、,.]+", "", purpose_reference)
                generic_references = {
                    "已确认",
                    "用户确认",
                    "用户已确认",
                    "用户说明",
                    "用户陈述",
                    "用户明确确认",
                }
                generic_reference_pattern = re.compile(
                    r"^(?:用户|委托人|当事人)?(?:在?本任务中|在?本对话中)?"
                    r"(?:已|明确)?(?:确认|说明|陈述)(?:证明)?(?:用途)?$"
                )
                has_speaker_anchor = bool(re.search(r"用户|委托人|当事人|本任务|本对话", purpose_reference))
                has_purpose_anchor = bool(
                    re.search(
                        r"请求|抗辩|主张|反驳|要件|返还|赔偿|履行|支付|解除|驳回|撤销|"
                        r"无效|确认(?:合同|关系|权利|义务|事实)",
                        purpose_reference,
                    )
                )
                if (
                    compact_reference in generic_references
                    or generic_reference_pattern.fullmatch(compact_reference)
                    or len(compact_reference) < 10
                    or not has_speaker_anchor
                    or not has_purpose_anchor
                ):
                    blockers.append(
                        f"{label}.purpose_basis.reference 未形成可复核的用户说明留痕；"
                        "应记录用户在本任务/本对话中明确确认的具体请求、抗辩或证明用途"
                    )
            if strict_v2:
                target_kind = purpose_basis.get("target_kind")
                target_id = purpose_basis.get("target_id")
                legal_element = purpose_basis.get("legal_element")
                if target_kind not in PURPOSE_TARGET_KINDS:
                    blockers.append(
                        f"{label}.purpose_basis.target_kind 无效；"
                        "只能是 claim、defense、rebuttal 或 context"
                    )
                if not isinstance(target_id, str) or not target_id.strip():
                    blockers.append(f"{label}.purpose_basis.target_id 必须回指具体请求、抗辩或背景目标")
                if not isinstance(legal_element, str) or not legal_element.strip():
                    blockers.append(f"{label}.purpose_basis.legal_element 必须写明对应法律要件或明确背景用途")
                if isinstance(target_id, str) and target_id.strip() and target_kind in PURPOSE_TARGET_KINDS:
                    target_key = (_text(target_kind), target_id.strip())
                    if target_key not in target_registry:
                        blockers.append(
                            f"{label}.purpose_basis 指向未声明的 proof_target：{target_key[0]}/{target_key[1]}"
                        )
                    elif isinstance(legal_element, str) and legal_element.strip() not in target_registry[target_key]:
                        blockers.append(
                            f"{label}.purpose_basis.legal_element 未列入 {target_key[0]}/{target_key[1]}："
                            f"{legal_element.strip()}"
                        )
        refs = _check_refs(claim.get("material_refs"), material_ids, label, blockers, strict_v2=strict_v2)
        claim["material_refs"] = refs
        ref_ids = {_text(ref.get("material_id")) for ref in refs if isinstance(ref, dict) and _text(ref.get("material_id"))}
        ref_groups = {material_group_by_id.get(material_id) for material_id in ref_ids if material_group_by_id.get(material_id)}
        container_group = _text(claim.get("_container_group_id"))
        declared_group = _text(claim.get("group_id"))
        explicit_cross_marker = any(
            claim.get(key) is True
            or _text(claim.get(key)).lower() in {"cross_group", "cross-group", "跨组", "true", "yes"}
            for key in ("cross_group", "allow_cross_group", "group_scope", "scope", "mode", "evidence_mode")
        )
        if _text(declared_group).lower() in {"cross_group", "cross-group", "跨组"}:
            explicit_cross_marker = True
            declared_group = ""
            claim["_cross_group"] = True
        if declared_group not in group_ids and declared_group in group_number_aliases:
            declared_group = group_number_aliases[declared_group]
            claim["_normalised_group_id"] = declared_group
        owner_group = container_group or declared_group
        if container_group and declared_group and declared_group != container_group:
            blockers.append(f"{label} 的 group_id={declared_group} 与所在证据组 {container_group} 不一致")
        if not container_group and declared_group and declared_group not in group_ids:
            blockers.append(f"{label} 指向不存在的 group_id：{declared_group}")
            owner_group = ""
        if not owner_group:
            if len(ref_groups) == 1:
                owner_group = next(iter(ref_groups))
            elif len(ref_groups) > 1:
                owner_group = sorted(ref_groups, key=lambda gid: group_order.get(gid, 10**9))[0]
                claim["_cross_group"] = True
            else:
                blockers.append(f"{label} 无法从 material_refs 确定唯一证据组归属")
        if owner_group and ref_groups and (ref_groups - {owner_group}):
            claim["_cross_group"] = True
            claim["_cross_group_groups"] = sorted(ref_groups, key=lambda gid: group_order.get(gid, 10**9))
        if claim.get("_cross_group"):
            if not explicit_cross_marker:
                blockers.append(f"{label} 的 material_refs 跨证据组，必须显式标记 cross_group=true")
            if strict_v2 and role != "linking":
                blockers.append(f"{label} 跨证据组共同证明时 role 必须为 linking")
            cross_group_ids = list(claim.get("_cross_group_groups", []))
            purpose_text_for_groups = _text(purpose)
            names_every_group = all(
                gid in purpose_text_for_groups
                or bool(
                    group_number_by_id.get(gid)
                    and re.search(
                        rf"第\s*{re.escape(group_number_by_id[gid])}\s*(?:组|项)",
                        purpose_text_for_groups,
                    )
                )
                for gid in cross_group_ids
            )
            if strict_v2 and (
                not re.search(r"结合|衔接|共同|交叉|相互印证|证据链", purpose_text_for_groups)
                or not cross_group_ids
                or not names_every_group
            ):
                blockers.append(f"{label} 跨证据组共同证明时 purpose 必须明确说明组间衔接")
        elif strict_v2 and role == "linking":
            blockers.append(f"{label} 使用 role=linking 时必须实际跨组回指并显式标记 cross_group=true")
        claim["_render_group_id"] = owner_group
        for field in ("text", "atomic_fact", "fact", "legal_significance", "boundary", "purpose"):
            if _nonempty(claim.get(field)):
                blockers.extend(_risk_messages(_text(claim.get(field)), f"{label}.{field}"))
        if isinstance(purpose_basis, dict) and _nonempty(purpose_basis.get("reference")):
            blockers.extend(
                _risk_messages(
                    _text(purpose_basis.get("reference")),
                    f"{label}.purpose_basis.reference",
                )
            )
    if normalized.get("_source_schema_version") == 2 and not claims:
        blockers.append("v2 必须提供非空 proof_claims[]，并逐项回指 material_refs")

    for group in groups:
        if not isinstance(group, dict):
            continue
        group_label = f"第{group.get('number', '?')}组"
        if _nonempty(group.get("proof_object")):
            blockers.extend(_date_completeness_messages(_text(group.get("proof_object")), f"{group_label} proof_object"))
            if strict_v2 and _proof_object_is_material_summary(group.get("proof_object")):
                blockers.append(
                    f"{group_label} proof_object 仅概括材料名称或内容；"
                    "必须改写为具体待证事实及其对应请求、抗辩、法律要件或明确背景用途"
                )
            if strict_v2 and not _proof_object_has_target_anchor(group.get("proof_object")):
                blockers.append(
                    f"{group_label} proof_object 未说明该材料服务的具体请求、抗辩、法律要件或明确背景用途"
                )
        materials = group.get("materials", [])
        if isinstance(materials, list):
            for material in materials:
                if isinstance(material, dict) and _nonempty(material.get("source_time")):
                    blockers.extend(
                        _date_completeness_messages(
                            _text(material.get("source_time")),
                            f"材料 {_text(material.get('material_id'))}.source_time",
                        )
                    )
    for claim in claims:
        if isinstance(claim, dict) and _nonempty(claim.get("text")):
            blockers.extend(_date_completeness_messages(_text(claim.get("text")), f"proof_claim {_text(claim.get('claim_id'))}.text"))

    # Ensure every valid claim has one, and only one, render owner.  The
    # owner is computed above so a top-level cross-group claim is displayed
    # once with a cross-group marker instead of being duplicated in two rows.
    render_counts = {claim_id: 0 for claim_id in claim_ids if claim_id}
    for claim in claims:
        if isinstance(claim, dict) and _text(claim.get("_render_group_id")) in group_ids:
            render_counts[_text(claim.get("claim_id"))] = render_counts.get(_text(claim.get("claim_id")), 0) + 1
    for claim_id, count in render_counts.items():
        if count != 1:
            blockers.append(f"proof_claim {claim_id} 必须恰好归属并渲染一次，当前归属数={count}")

    if strict_v2:
        claims_per_group = {group_id: 0 for group_id in group_ids}
        for claim in claims:
            if not isinstance(claim, dict):
                continue
            covered_group_ids = {_text(claim.get("_render_group_id"))}
            covered_group_ids.update(_text(item) for item in _as_list(claim.get("_cross_group_groups")))
            for covered_group_id in covered_group_ids:
                if covered_group_id in claims_per_group:
                    claims_per_group[covered_group_id] += 1
        for group_id, count in claims_per_group.items():
            if count == 0:
                blockers.append(f"证据组 {group_id} 没有可回指的 proof_claim，材料尚未形成明确证明用途")

    coverage_matrix = build_proof_coverage_matrix(normalized)
    normalized["_coverage_matrix"] = coverage_matrix
    if strict_v2:
        for material_id in coverage_matrix.get("materials_without_purpose", []):
            blockers.append(f"材料 {material_id} 未被任何具有 purpose 的 proof_claim 回指；属于有材料但无证明用途")
        for claim_id in coverage_matrix.get("claims_without_target", []):
            blockers.append(f"proof_claim {claim_id} 未映射到具体请求/抗辩/法律要件或明确背景用途")
        for row in coverage_matrix.get("coverage", []):
            if isinstance(row, dict) and row.get("status") == "uncovered":
                blockers.append(
                    "法律要件缺少证据覆盖："
                    f"{_text(row.get('target_kind'))}/{_text(row.get('target_id'))}/"
                    f"{_text(row.get('legal_element'))}"
                )
            elif isinstance(row, dict) and row.get("status") == "partial":
                warnings.append(
                    "覆盖矩阵提示仅有间接、辅助或补强材料，需人工复核充分性："
                    f"{_text(row.get('target_kind'))}/{_text(row.get('target_id'))}/"
                    f"{_text(row.get('legal_element'))}"
                )

    for group in groups:
        if not isinstance(group, dict):
            continue
        blockers.extend(_risk_messages(_text(group.get("proof_object")), f"第{group.get('number', '?')}组 proof_object"))
        for key in ("atomic_fact", "fact", "legal_significance", "boundary"):
            if _nonempty(group.get(key)):
                blockers.extend(_risk_messages(_text(group.get(key)), f"第{group.get('number', '?')}组 {key}"))
        for material in group.get("materials", []) if isinstance(group.get("materials"), list) else []:
            if isinstance(material, dict):
                for key in ("fact", "atomic_fact", "legal_significance", "boundary"):
                    if _nonempty(material.get(key)):
                        blockers.extend(_risk_messages(_text(material.get(key)), f"材料 {_text(material.get('material_id'))}.{key}"))
    for key, check in validation.get("checks", {}).items() if isinstance(validation.get("checks"), dict) else []:
        if isinstance(check, dict) and _nonempty(check.get("note")):
            blockers.extend(_risk_messages(_text(check.get("note")), f"核验项 {key}.note"))

    material_root_values = {str(path.parent.resolve()) for path in material_paths}
    material_common_root = ""
    if material_paths:
        try:
            common_root = Path(os.path.commonpath([str(path.resolve()) for path in material_paths]))
            if common_root.is_file():
                common_root = common_root.parent
            material_common_root = str(common_root.resolve())
        except (OSError, ValueError):
            material_common_root = ""
    # A common path can be a nested subdirectory (for example
    # ``materials/chat/M001.png``).  Protect an explicitly declared material
    # root and the top-level material tree under the manifest directory too,
    # so an output cannot escape through ``materials/bundle.docx`` merely
    # because every input happens to live below ``materials/chat``.
    declared_root_values: set[str] = set()
    for root_key in ("material_root", "materials_root", "material_dir", "materials_dir"):
        raw_root = normalized.get(root_key)
        if not isinstance(raw_root, str) or not raw_root.strip():
            continue
        try:
            declared_root_values.add(str(resolve_material_path(raw_root, base_dir)))
        except ValueError as exc:
            blockers.append(f"{root_key} 无效：{exc}")
    base_resolved = base_dir.resolve()
    if material_common_root:
        common_path = Path(material_common_root)
        try:
            relative_common = common_path.relative_to(base_resolved)
        except ValueError:
            relative_common = None
        if relative_common is not None and relative_common.parts:
            top_level = base_resolved / relative_common.parts[0]
            if top_level.is_dir():
                declared_root_values.add(str(top_level.resolve()))
        # Recognise conventional material-root names for absolute manifests;
        # an explicitly supplied ``material_root`` remains the authoritative
        # option for projects using a different name.
        root_names = {"materials", "material", "evidence", "case-materials", "证据材料", "附件"}
        for ancestor in (common_path, *common_path.parents):
            if ancestor == base_resolved:
                break
            if ancestor.name.lower() in root_names:
                declared_root_values.add(str(ancestor.resolve()))
    material_root_values.update(declared_root_values)
    normalized["_preflight"] = {
        "records": material_records,
        "duplicate_groups": duplicate_groups,
        "alias_mapping": alias_mapping,
        "material_paths": material_paths,
        "material_roots": sorted(material_root_values),
        "material_common_root": material_common_root,
        "material_ids": sorted(material_ids),
        "context_material_ids": sorted(context_ids),
        "warnings": warnings,
        "blockers": blockers,
    }
    return normalized, fatal, blockers, warnings


def _output_path_errors(normalized: dict[str, Any], output_path: Path, manifest_path: Path, overwrite: bool) -> list[str]:
    errors: list[str] = []
    raw_output_path = output_path
    output_path = output_path.resolve()
    if output_path == manifest_path.resolve():
        errors.append(f"输出路径不得等于 manifest：{output_path}")
    if output_path.exists() and output_path.is_dir():
        errors.append(f"输出路径必须是文件，不能覆盖目录：{output_path}")
    link_component = first_link_component(raw_output_path)
    if link_component is not None:
        errors.append(f"拒绝将输出写入符号链接或目录联接：{link_component}")
    for material_path in normalized.get("_preflight", {}).get("material_paths", []):
        if output_path == material_path.resolve():
            errors.append(f"输出路径不得覆盖输入材料：{output_path}")
    roots_to_protect = list(normalized.get("_preflight", {}).get("material_roots", []))
    common_root = normalized.get("_preflight", {}).get("material_common_root")
    if common_root:
        roots_to_protect.append(common_root)
    for root_text in roots_to_protect:
        root = Path(root_text).resolve()
        try:
            output_path.relative_to(root)
        except ValueError:
            continue
        errors.append(f"输出路径不得落入材料目录（只读保护）：{output_path}；材料目录：{root}")
        break
    if output_path.exists() and not overwrite:
        errors.append(f"拒绝覆盖已存在输出（默认保护）：{output_path}；如确需覆盖请显式使用 --overwrite")
    return errors


def validate_manifest(
    data: dict[str, Any],
    base_dir: Path,
    allow_draft: bool = False,
    output_path: Path | None = None,
    manifest_path: Path | None = None,
    overwrite: bool = False,
) -> dict[str, Any]:
    """Validate and annotate a manifest; return a build-ready copy.

    Fatal structural/input errors always stop. Formal blockers may be
    rendered only with ``allow_draft``; v1 is automatically forced to draft.
    """

    normalized, fatal, blockers, warnings = _prepare(data, base_dir, manifest_path)
    if output_path is not None and manifest_path is not None:
        fatal.extend(_output_path_errors(normalized, output_path, manifest_path, overwrite))
    if fatal:
        raise ManifestError(fatal, warnings)
    # v1 is readable but never silently becomes a formal document. The caller
    # must explicitly request --allow-draft, even when legacy booleans are true.
    if blockers and not (allow_draft or normalized.get("_intake_draft_authorized")):
        raise ManifestError(blockers, warnings)
    normalized["_effective_draft"] = bool(allow_draft or normalized.get("_forced_draft"))
    normalized["_warnings"] = warnings
    normalized["_blockers"] = blockers
    return normalized


def set_run_font(run, name="宋体", size=12, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:color"), "000000")


def set_table_column_widths(table, widths: list[Any]) -> None:
    """Write fixed column widths to every OOXML location Word/LibreOffice reads.

    ``python-docx``'s ``cell.width`` alone does not reliably update the table
    grid used by LibreOffice.  Keep the public width list in centimetres but
    persist its twip values in ``tblGrid``, table layout/width and every cell's
    ``tcW`` so the five-column filing table keeps its intended proportions.
    """

    if not widths:
        return
    twips = [max(1, int(round(width.inches * 1440))) for width in widths]
    table_width = sum(twips)
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(table_width))

    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        insert_at = 1 if len(tbl) and tbl[0].tag == qn("w:tblPr") else 0
        tbl.insert(insert_at, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width_twips in twips:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width_twips))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index >= len(twips):
                continue
            cell.width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(twips[index]))


def assert_table_column_widths(table, widths: list[Any]) -> None:
    """Minimal OOXML regression assertion; do not rely on ``cell.width``."""

    expected = [max(1, int(round(width.inches * 1440))) for width in widths]
    tbl_grid = table._tbl.find(qn("w:tblGrid"))
    actual = []
    if tbl_grid is not None:
        actual = [int(child.get(qn("w:w"))) for child in tbl_grid.findall(qn("w:gridCol"))]
    if actual != expected:
        raise RuntimeError(f"证据表 tblGrid 宽度未按声明写入：expected={expected} actual={actual}")
    if len(expected) == 5:
        proof_ratio = actual[-1] / max(1, sum(actual))
        if not 0.45 <= proof_ratio <= 0.50:
            raise RuntimeError(f"证据表证明对象列宽度比例异常：{proof_ratio:.4f}")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index >= len(expected):
                continue
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None or int(tc_w.get(qn("w:w"))) != expected[index]:
                raise RuntimeError(f"证据表第{index + 1}列 tcW 未同步写入")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_text(paragraph, text, size=12, bold=False, align=None, color=None):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(_text(text))
    set_run_font(run, size=size, bold=bold)
    if color:
        run.font.color.rgb = color
    return run


def add_word_field(paragraph, instruction: str, fallback: str, size: int = 10):
    """Add a complex Word field with a readable pre-update result."""

    begin_run = paragraph.add_run()
    set_run_font(begin_run, size=size)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    set_run_font(instruction_run, size=size)
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = f" {instruction.strip()} "
    instruction_run._r.append(instruction_text)

    separate_run = paragraph.add_run()
    set_run_font(separate_run, size=size)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    result_run = paragraph.add_run(_text(fallback))
    set_run_font(result_run, size=size)

    end_run = paragraph.add_run()
    set_run_font(end_run, size=size)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    return result_run


def set_update_fields_on_open(document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_page_number_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_text(paragraph, "第", size=10)
    add_word_field(paragraph, "PAGE", "1", size=10)
    add_text(paragraph, "页", size=10)


def attachment_bookmark_name(attachment_page: int) -> str:
    return f"evidence_page_{attachment_page:04d}"


def add_paragraph_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_page_range(paragraph, attachment_pages: list[int]) -> None:
    """Render group page locations with PAGEREF fields.

    Each attachment occupies one physical page.  Word/LibreOffice resolves the
    bookmarks after pagination, so a multi-page evidence table does not shift
    hard-coded ranges.  The fallback assumes a one-page index until fields are
    refreshed on open.
    """

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    pages = sorted(set(page for page in attachment_pages if page > 0))
    if not pages:
        add_text(paragraph, "—", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        return

    add_text(paragraph, "第", size=10)
    contiguous = pages == list(range(pages[0], pages[-1] + 1))
    display_pages = [pages[0], pages[-1]] if contiguous and len(pages) > 1 else pages
    for index, page in enumerate(display_pages):
        if index:
            add_text(paragraph, "-" if contiguous else "、", size=10)
        add_word_field(
            paragraph,
            f"PAGEREF {attachment_bookmark_name(page)}",
            str(page + 1),
            size=10,
        )
    add_text(paragraph, "页", size=10)


def _normalised_image_bytes(
    path: Path,
    source_digest: str | None = None,
    material_id: str | None = None,
    display_rotation_degrees_clockwise: int = 0,
) -> tuple[io.BytesIO, tuple[int, int]]:
    from PIL import Image, ImageOps

    display_rotation_degrees_clockwise = _display_rotation(display_rotation_degrees_clockwise)
    if source_digest is not None and not re.fullmatch(r"[0-9a-fA-F]{64}", _text(source_digest)):
        raise ValueError("source_sha256 必须是原文件的64位 SHA-256，不得拼接材料编号或其他文本")
    with Image.open(path) as image:
        image.load()
        normalized = ImageOps.exif_transpose(image)
        normalized.load()
        if display_rotation_degrees_clockwise:
            normalized = normalized.rotate(-display_rotation_degrees_clockwise, expand=True)
            normalized.load()
        if normalized.mode not in {"RGB", "RGBA"}:
            normalized = normalized.convert("RGBA" if "transparency" in normalized.info else "RGB")
        # The embedded image is a阅卷副本.  Do not carry source EXIF/ICC or
        # other image metadata into the DOCX; provenance remains in Manifest.
        normalized.info.clear()
        buffer = io.BytesIO()
        normalized.save(buffer, format="PNG", optimize=False)
        buffer.seek(0)
        return buffer, normalized.size


def fit_picture(
    run,
    path,
    max_width_cm,
    max_height_cm,
    source_digest: str | None = None,
    material_id: str | None = None,
    display_rotation_degrees_clockwise: int = 0,
):
    buffer, (width_px, height_px) = _normalised_image_bytes(
        Path(path),
        source_digest=source_digest,
        material_id=material_id,
        display_rotation_degrees_clockwise=display_rotation_degrees_clockwise,
    )
    ratio = width_px / max(1, height_px)
    max_ratio = max_width_cm / max_height_cm
    if ratio >= max_ratio:
        return run.add_picture(buffer, width=Cm(max_width_cm))
    return run.add_picture(buffer, height=Cm(max_height_cm))


def _attachment_image_height_cm(header_text: str) -> float:
    """Reserve enough page height for a wrapped attachment header.

    ``keep_with_next`` keeps the header and image together.  Estimate wrapped
    8-point Chinese lines conservatively and fit the image into the remaining
    A4 content height; full provenance remains in the Manifest/QA sidecar.
    """

    content_height_cm = 29.7 - 2.54 - 2.54
    chars_per_line = 46
    line_count = max(1, math.ceil(len(header_text) / chars_per_line))
    header_height_cm = 0.38 * line_count + 0.18
    # Leave a small cushion for paragraph leading and the image anchor.
    return max(6.0, min(22.0, content_height_cm - header_height_cm - 0.22))


def _all_material_items(data: dict[str, Any]) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    pairs: list[tuple[dict[str, Any], dict[str, Any]]] = []
    for group in data.get("evidence_groups", []):
        for material in group.get("materials", []):
            pairs.append((group, material))
    return pairs


def _plan_attachments(data: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    records = data.get("_preflight", {}).get("records", {})
    attachments: list[dict[str, Any]] = []
    index: dict[str, dict[str, Any]] = {}
    all_items = _all_material_items(data)
    items_by_id = {_text(material.get("material_id")): (group, material) for group, material in all_items if isinstance(material, dict)}
    canonical_by_id = {
        material_id: _text(records.get(material_id, {}).get("canonical_material_id", material_id)) or material_id
        for material_id in items_by_id
    }
    # Embed the declared canonical material at its own manifest position, even
    # when an alias appears earlier.  The alias index is populated afterwards
    # with the canonical's actual attachment number.
    embedded_canonicals: set[str] = set()
    page_by_canonical: dict[str, int] = {}
    group_page_counts: dict[str, int] = {}
    for group, material in all_items:
        material_id = _text(material.get("material_id"))
        canonical = canonical_by_id.get(material_id, material_id)
        if material_id != canonical or canonical in embedded_canonicals:
            continue
        canonical_group, canonical_material = items_by_id.get(canonical, (group, material))
        record = records.get(canonical, records.get(material_id, {}))
        page = len(attachments) + 1
        canonical_group_id = _text(canonical_group.get("group_id"))
        group_page_counts[canonical_group_id] = group_page_counts.get(canonical_group_id, 0) + 1
        item = {
            "group": canonical_group,
            "material": canonical_material,
            "record": record,
            "page": page,
            "group_page": group_page_counts[canonical_group_id],
            "canonical": canonical,
            "embedded": True,
        }
        attachments.append(item)
        embedded_canonicals.add(canonical)
        page_by_canonical[canonical] = page
    # If a malformed declaration points outside the listed materials, preserve
    # a usable fallback item for draft inspection; formal validation already
    # reports the declaration error.
    for group, material in all_items:
        material_id = _text(material.get("material_id"))
        canonical = canonical_by_id.get(material_id, material_id)
        record = records.get(material_id, {})
        canonical_item = next(
            (item for item in attachments if item.get("canonical") == canonical),
            {},
        )
        index[material_id] = {
            "group": group,
            "material": material,
            "record": record,
            "page": page_by_canonical.get(canonical, ""),
            "group_page": canonical_item.get("group_page", ""),
            "canonical": canonical,
            "embedded": material_id == canonical and canonical in page_by_canonical,
        }
    return attachments, index


def _proof_text(data: dict[str, Any], group: dict[str, Any]) -> str:
    # Keep the five-column filing table concise.  Atomic claims and their
    # provenance remain in the Manifest/sidecar review materials, not in the
    # lawyer-facing filing volume.
    return _text(group.get("_display_proof_object") or _display_proof_object(group.get("proof_object")))


def _explicit_missing_categories(data: dict[str, Any]) -> list[str]:
    """Return privacy-safe categories from structured unresolved-issue notes."""

    validation = data.get("validation", {})
    issues = validation.get("unresolved_issues", []) if isinstance(validation, dict) else []
    if not isinstance(issues, list):
        return []
    category_rules = (
        (("起诉状", "答辩状", "仲裁申请", "仲裁请求", "代理意见", "书状", "请求与抗辩"), "起诉状、答辩状或仲裁请求等相关书状"),
        (("案件信息", "案号", "案由"), "案件基本信息"),
        (("提交方", "提交人", "签名", "提交时间"), "提交方信息"),
        (("主体", "身份", "昵称", "头像", "群名"), "主体身份或对应关系"),
        (("日期", "年份", "时间"), "日期或时间线"),
        (("金额", "申请额", "订单额", "实付额", "认可额", "结算额"), "金额口径"),
        (("付款", "支付", "清偿"), "付款或清偿状态"),
        (("上下文", "裁切", "前后文"), "关键上下文"),
        (("证明用途", "证明对象", "待证事实", "请求", "抗辩", "证明要件"), "证明用途及其与请求、抗辩的对应关系"),
        (("可读性", "低清", "损坏", "不可读", "模糊"), "材料可读性"),
        (("原始载体", "原件", "原应用", "回单", "录屏"), "原始载体或可回溯位置"),
        (("其他",), "其他待补信息"),
    )
    categories: list[str] = []
    for issue in issues:
        if not isinstance(issue, str) or not issue.strip():
            continue
        for tokens, label in category_rules:
            if any(token in issue for token in tokens) and label not in categories:
                categories.append(label)
    return categories


def _draft_issue_summary(data: dict[str, Any]) -> str:
    """Return a lawyer-facing, non-technical draft warning summary.

    Detailed blocker strings remain available to the Manifest/QA sidecar. A
    filing-facing DOCX must not expose paths, hashes, internal IDs, carrier
    field names, or implementation diagnostics.
    """

    blockers = data.get("_blockers", [])
    blocker_text = " ".join(_text(item) for item in blockers)
    missing_categories = _explicit_missing_categories(data)
    topics: list[str] = []
    if any(token in blocker_text for token in ("purpose", "用途", "请求、抗辩", "证明要件")):
        topics.append("部分证明对象尚未说明拟证明的具体事实，以及其服务的诉讼请求、仲裁请求、抗辩或证明要件")
    if any(token in blocker_text for token in ("主体", "identity", "身份")):
        topics.append("部分主体关系仍需结合完整资料核对")
    if any(token in blocker_text for token in ("日期", "dates", "时间")):
        topics.append("部分日期和时间线仍需核对")
    if any(token in blocker_text for token in ("金额", "付款", "amount", "payment")):
        topics.append("部分金额口径或付款状态仍需核对")
    if any(token in blocker_text for token in ("原始载体", "original_carrier", "载体")):
        topics.append("部分材料的原始载体和来源仍需补充定位")
    if not topics:
        topics.append("材料依据、证明范围或文书对应关系仍需人工复核")
    missing_summary = f"明确缺失：{'、'.join(missing_categories)}。" if missing_categories else ""
    return (
        "当前仅供内部核对，尚不具备正式举证条件。"
        + missing_summary
        + "；".join(topics)
        + "。用户选择不知道、不提供或稍后补充时仍可继续整理，但上述缺口不会被视为已经解决；"
        "请补充相关资料并逐项核对后再生成正式版本。"
    )


def build_document(data: dict[str, Any], output_path: Path, base_dir: Path, allow_draft: bool = False):
    office_fatal, office_blockers, office_warnings = validate_office_provenance(data, base_dir)
    if office_fatal or office_blockers:
        raise ManifestError(office_fatal + office_blockers, office_warnings)
    if "_preflight" not in data:
        data = validate_manifest(data, base_dir, allow_draft=allow_draft, output_path=output_path, overwrite=False)
    effective_draft = bool(allow_draft or data.get("_effective_draft"))
    attachments, index = _plan_attachments(data)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(3.175)
    section.footer_distance = Cm(1.27)
    set_update_fields_on_open(document)
    add_page_number_footer(section)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    title = document.add_paragraph()
    add_text(title, "证据清单", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_after = Pt(6)
    if effective_draft:
        warning = document.add_paragraph()
        add_text(warning, "内部核对稿", size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    display_case_info = _display_case_info(data.get("case_info", ""))
    if display_case_info:
        add_text(document.add_paragraph(), f"案件信息：{display_case_info}", bold=True)
    add_text(document.add_paragraph(), f"提交方：{data.get('submitter', '')}", bold=True)
    add_text(document.add_paragraph(), f"提交人签名：{data.get('signer', '')}                提交时间：{data.get('submission_date', '')}", bold=True)

    groups = data["evidence_groups"]
    table = document.add_table(rows=1, cols=5)
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(0.78), Cm(3.43), Cm(2.10), Cm(1.64), Cm(6.70)]
    headers = ["编号", "证据名称", "页码", "证据形式", "证明对象"]
    set_repeat_table_header(table.rows[0])
    for index_number, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[index_number]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        add_text(cell.paragraphs[0], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    # A group containing only an alias still points to the canonical attachment
    # bookmark, so its table page reference remains usable without embedding a
    # duplicate image.
    pages_by_group: dict[str, set[int]] = {}
    for group, material in _all_material_items(data):
        material_id = _text(material.get("material_id"))
        page_value = index.get(material_id, {}).get("page")
        if isinstance(page_value, int) and page_value > 0:
            pages_by_group.setdefault(_text(group.get("group_id")), set()).add(page_value)
    for group in groups:
        row = table.add_row()
        gid = _text(group.get("group_id"))
        values = [
            str(group.get("number", "")),
            _text(group.get("evidence_name")),
            "",
            _text(group.get("evidence_form")),
            _proof_text(data, group),
        ]
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            if idx == 2:
                add_page_range(cell.paragraphs[0], sorted(pages_by_group.get(gid, set())))
            else:
                add_text(cell.paragraphs[0], value, size=10 if idx == 4 else 12, align=align)
    set_table_column_widths(table, widths)
    assert_table_column_widths(table, widths)

    for attachment_index, attachment in enumerate(attachments):
        group = attachment["group"]
        material = attachment["material"]
        record = attachment["record"]
        header = document.add_paragraph()
        header.paragraph_format.page_break_before = True
        header.paragraph_format.keep_with_next = True
        header_text = (
            f"证据编号：{group.get('number', '?')}｜"
            f"证据名称：{_text(group.get('evidence_name'))}｜"
            f"组内页次：{attachment.get('group_page', attachment['page'])}"
        )
        add_text(header, header_text, size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
        add_paragraph_bookmark(
            header,
            attachment_bookmark_name(int(attachment["page"])),
            bookmark_id=attachment_index + 1,
        )
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run()
        source_digest = record.get("sha256")
        fit_picture(
            run,
            record["path"],
            max_width_cm=14.65,
            max_height_cm=_attachment_image_height_cm(header_text),
            source_digest=source_digest,
            material_id=_text(material.get("material_id")),
            display_rotation_degrees_clockwise=int(record.get("display_rotation_degrees_clockwise", 0)),
        )

    core = document.core_properties
    core.title = "证据清单及证据材料"
    core.subject = ""
    core.author = ""
    core.last_modified_by = ""
    core.keywords = ""
    core.comments = ""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def _load_json(path: Path) -> dict[str, Any]:
    try:
        loaded = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ManifestError([f"读取 manifest 失败：{exc}"]) from None
    if not isinstance(loaded, dict):
        raise ManifestError(["manifest 顶层必须是 JSON 对象"])
    return loaded


def _write_coverage_matrix(data: dict[str, Any], output_path: Path) -> None:
    matrix = data.get("_coverage_matrix")
    if not isinstance(matrix, dict):
        matrix = build_proof_coverage_matrix(data)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(matrix, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="生成证据清单＋内嵌证据材料 DOCX")
    parser.add_argument("manifest", help="Manifest JSON 路径")
    parser.add_argument("output", help="输出 DOCX 路径")
    parser.add_argument("--allow-draft", action="store_true", help="允许将存在形式阻断的 v2 清单输出为草稿")
    parser.add_argument("--overwrite", "--force", dest="overwrite", action="store_true", help="显式允许覆盖已有输出")
    parser.add_argument(
        "--coverage-matrix",
        help="内部覆盖矩阵 JSON 路径；默认与 DOCX 同目录并使用 .coverage-matrix.json 后缀",
    )
    args = parser.parse_args()
    manifest_path = Path(args.manifest).expanduser().resolve()
    output_path = Path(args.output).expanduser()
    if not output_path.is_absolute():
        output_path = Path.cwd() / output_path
    coverage_output = (
        Path(args.coverage_matrix).expanduser()
        if args.coverage_matrix
        else output_path.with_name(output_path.stem + ".coverage-matrix.json")
    )
    if not coverage_output.is_absolute():
        coverage_output = Path.cwd() / coverage_output
    try:
        data = _load_json(manifest_path)
        normalized = validate_manifest(data, manifest_path.parent, allow_draft=args.allow_draft, output_path=output_path, manifest_path=manifest_path, overwrite=args.overwrite)
        coverage_errors = _output_path_errors(
            normalized,
            coverage_output,
            manifest_path,
            args.overwrite,
        )
        if coverage_output.resolve() == output_path.resolve():
            coverage_errors.append("覆盖矩阵输出不得与 DOCX 输出为同一文件")
        if coverage_errors:
            raise ManifestError(coverage_errors)
        build_document(normalized, output_path, manifest_path.parent, allow_draft=args.allow_draft)
        _write_coverage_matrix(normalized, coverage_output)
    except ManifestError as error:
        raise SystemExit(f"生成失败：\n{error}") from None
    except (OSError, ValueError, RuntimeError) as error:
        raise SystemExit(f"生成失败：\n- {error}") from None
    print(f"已生成：{output_path}")
    print(f"覆盖矩阵：{coverage_output}")
    print(f"证据组：{len(normalized['evidence_groups'])}；实际嵌入材料页：{len(_plan_attachments(normalized)[0])}")
    if normalized.get("_effective_draft"):
        print("提示：本文件为内部核对稿。")
        print(f"交付回复须说明的缺失/待核事项：{_draft_issue_summary(normalized)}")
    for warning in normalized.get("_warnings", []):
        print(f"警示：{warning}")


if __name__ == "__main__":
    main()
